"""Ҳуҷҷатнигорӣ → Ҳисоботи синфҳо — quarterly class report Excel export.

Reproduces the owner's «Ҳисобот» worksheet (portal/excel_templates/
class_report.xlsx) cell-for-cell: same columns, header block, hidden
Таблица1 header row, per-student band formulas, footer summary rows and
signature lines. Only the ranges/row positions adapt to the class size —
the owner's formula structure is preserved verbatim.

Data sources (existing portal data only):
  - Student (school, class_name, full_name, gender)      -> rows, C, D
  - ClassSubject (is_active, minus UNGRADED_SUBJECTS)     -> subject columns E..X
  - QuarterGrade.grade for the chosen quarter             -> subject cells
  - СОАТИ ТАРБИЯВӢ ClassSubject allocated_teacher/teacher -> роҳбари синф (AG1)
  - School.name / School.director                         -> title + signature
"""
import datetime
import io
import re
from copy import copy
from pathlib import Path

from django.conf import settings
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from .curriculum_hours import UNGRADED_SUBJECTS
from .models import (
    ClassSubject, QuarterGrade, Student, TeacherProfile,
    normalize_class_name, normalize_subject,
)
from .utils import default_subjects_for_class

TEMPLATE_PATH = (
    Path(settings.BASE_DIR) / 'portal' / 'excel_templates' / 'class_report.xlsx'
)

MAX_COL = 36            # A..AJ
FIRST_DATA_ROW = 6      # first student row in the template
TEMPLATE_DATA_END = 38  # last student row in the template (33 slots)
SUBJ_FIRST_COL = 5      # E
SUBJ_LAST_COL = 24      # X — owner's band formulas cover $E:$X (20 slots)
SUBJ_AVG_LAST = 'W'     # owner's AG/AH/AI formulas average E..W (X is spare)
TJB_SUBJECT_LIMIT = SUBJ_LAST_COL - SUBJ_FIRST_COL + 1

# Template footer rows (band block S..S+3, then Холи миёна / Фоиз /
# Шумораи хонандагон, spacer, signature, spacer, signature).
FOOTER_FIRST_ROW = 39
FOOTER_LAST_ROW = 49
ROW_HEIGHT = 12.5

QUARTER_ORDINALS = {1: '1-ум', 2: '2-юм', 3: '3-юм', 4: '4-ум'}


class ClassReportError(Exception):
    """Raised when a class cannot be rendered into the report template."""


# ---------------------------------------------------------------------------
# Data collection
# ---------------------------------------------------------------------------

def report_subjects(school, class_name):
    """Ordered graded subjects for the class: national-curriculum order
    first, then any school/registry extras (alphabetical)."""
    class_name = normalize_class_name(class_name)
    active = set(
        ClassSubject.objects.filter(
            school=school, class_name=class_name, is_active=True,
        ).values_list('subject', flat=True)
    ) - UNGRADED_SUBJECTS
    ordered = [
        s for s in (normalize_subject(x)
                    for x in default_subjects_for_class(class_name))
        if s in active
    ]
    ordered += sorted(active - set(ordered))
    return ordered


def class_teacher_name(school, class_name):
    """Роҳбари синф — the teacher allocated to СОАТИ ТАРБИЯВӢ."""
    cs = (
        ClassSubject.objects.filter(
            school=school, class_name=normalize_class_name(class_name),
            subject='СОАТИ ТАРБИЯВӢ', is_active=True,
        ).select_related('allocated_teacher', 'teacher').first()
    )
    if cs is None:
        return ''
    profile = cs.allocated_teacher
    if profile is None and cs.teacher_id:
        profile = TeacherProfile.objects.filter(user_id=cs.teacher_id).first()
    return profile.full_name if profile else ''


def _academic_year(today=None):
    """'2026-2027' for Sep 2026..Aug 2027."""
    today = today or datetime.date.today()
    year = today.year if today.month >= 9 else today.year - 1
    return f'{year}-{year + 1}'


def _school_short_name(school):
    """'Муассисаи таҳсилоти умумии №20' -> 'МТМУ №20' (owner's signature style)."""
    m = re.search(r'№\s*(\d+)', school.name or '')
    if m:
        return f'МТМУ №{m.group(1)}'
    return school.name


def _director_name(school):
    director = (school.director or '').strip()
    return '' if director in ('', 'Номи директор') else director


# ---------------------------------------------------------------------------
# Workbook construction
# ---------------------------------------------------------------------------

def _capture_row_style(ws, row):
    return {col: copy(ws.cell(row=row, column=col)._style)
            for col in range(1, MAX_COL + 1)}


def _capture_footer(ws):
    """(style, value) snapshot of every footer cell, rows 39..49."""
    spec = {}
    for row in range(FOOTER_FIRST_ROW, FOOTER_LAST_ROW + 1):
        spec[row] = {
            col: (copy(c._style), c.value)
            for col, c in enumerate(ws[row], start=1)
        }
    return spec


def build_class_report_workbook(school, class_name, quarter, today=None):
    """Build the «Ҳисобот» workbook for one school + class + quarter.

    Returns a BytesIO xlsx. Raises ClassReportError for an empty class or
    a subject list that cannot fit the template's 20 subject columns.
    """
    class_name = normalize_class_name(class_name)
    students = list(
        Student.objects.filter(school=school, class_name=class_name)
        .order_by('full_name')
    )
    if not students:
        raise ClassReportError('Дар синфи интихобшуда хонанда нест.')

    subjects = report_subjects(school, class_name)
    if len(subjects) > TJB_SUBJECT_LIMIT:
        raise ClassReportError(
            f'Синфи {class_name} {len(subjects)} фани фаъол дорад — '
            f'шаблони ҳисобот на зиёд аз {TJB_SUBJECT_LIMIT} фанро ҷой медиҳад.'
        )

    grades = {
        (sid, subj): grade
        for sid, subj, grade in QuarterGrade.objects.filter(
            student__in=students, class_name=class_name, quarter=quarter,
            grade__isnull=False,
        ).values_list('student_id', 'subject', 'grade')
    }
    teacher_name = class_teacher_name(school, class_name)

    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.worksheets[0]
    ws.title = class_name

    # Broken external-workbook leftovers (Baholar -> [1]Лист...): dropping
    # them keeps Excel from showing an "external links" warning.
    for name in ('Baholar', 'Fanlar'):
        if name in wb.defined_names:
            del wb.defined_names[name]

    n = len(students)
    last = FIRST_DATA_ROW + n - 1      # last student row
    s_row = last + 1                   # first footer band row (was 39)
    delta = s_row - FOOTER_FIRST_ROW

    data_styles = _capture_row_style(ws, FIRST_DATA_ROW)
    footer_spec = _capture_footer(ws)
    footer_merges = [
        m for m in list(ws.merged_cells.ranges) if m.min_row >= FOOTER_FIRST_ROW
    ]
    for m in footer_merges:
        ws.unmerge_cells(str(m))
    ws.delete_rows(FIRST_DATA_ROW, ws.max_row - FIRST_DATA_ROW + 1)

    _fill_header(ws, school, class_name, quarter, subjects, teacher_name, today)
    _fill_students(ws, students, subjects, grades, data_styles)
    _fill_footer(ws, footer_spec, footer_merges, delta, last, school)
    _resize_table(ws, last)
    _page_setup(ws, FOOTER_LAST_ROW + delta)

    wb.calculation.fullCalcOnLoad = True
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _fill_header(ws, school, class_name, quarter, subjects, teacher_name,
                 today):
    """Rows 1-5: title line, broken-INDIRECT fix, subject headers."""
    ordinal = QUARTER_ORDINALS.get(quarter, f'{quarter}-ум')
    school_text = (school.name or '').lower()
    ws['A2'] = (
        f'=CONCATENATE("чоряки {ordinal} дар синфи ", B6, " {school_text}'
        f' соли таҳсили {_academic_year(today)}")'
    )
    # Template A3 holds =INDIRECT("'"&$A$1&"'!A3") which evaluates to #REF!
    # (no 'Ҳисобот' sheet exists) — write the row-number header instead.
    ws['A3'] = '№'
    ws['AG1'] = teacher_name
    # Subject headers E3..X3 follow the class's own active subjects; unused
    # slots stay empty like the template's spare column X.
    for col in range(SUBJ_FIRST_COL, SUBJ_LAST_COL + 1):
        idx = col - SUBJ_FIRST_COL
        ws.cell(row=3, column=col).value = (
            subjects[idx] if idx < len(subjects) else None
        )
    ws['B6'] = class_name  # referenced by the A2 title formula


def _row_formulas(row, avg_last):
    """Owner's per-student formulas (Y..AI) for one data row."""
    ex = f'$E{row}:$X{row}'
    ew = f'E{row}:{avg_last}{row}'
    return {
        25: f'=IF(COUNT({ex})>0,COUNTIF({ex},">=9"),"")',                     # Y
        26: f'=IF(COUNT({ex})>0,Y{row}/COUNT({ex})*100,"")',                 # Z
        27: f'=IF(COUNT({ex})>0,COUNTIF({ex},">=7")-Y{row},"")',             # AA
        28: f'=IF(COUNT({ex})>0,AA{row}/COUNT({ex})*100,"")',                # AB
        29: f'=IF(COUNT({ex})>0,COUNTIF({ex},">=4")-(Y{row}+AA{row}),"")',   # AC
        30: f'=IF(COUNT({ex})>0,AC{row}/COUNT({ex})*100,"")',                # AD
        31: f'=IF(COUNT({ex})>0,COUNTIF({ex},"<4"),"")',                     # AE
        32: f'=IF(COUNT({ex})>0,AE{row}/COUNT({ex})*100,"")',                # AF
        33: f'=IF(COUNT({ew})>0,AVERAGE({ew}),"")',                          # AG
        34: f'=IF(COUNT({ew})>0,100,"")',                                    # AH
        35: f'=IF(COUNT({ew})>0,Z{row}+AB{row},"")',                         # AI
    }


def _fill_students(ws, students, subjects, grades, data_styles):
    # The owner's AG/AH/AI formulas average E:W (X is a spare slot). When a
    # class uses all 20 columns the range extends to X so every graded
    # subject is counted — identical results while X stays empty.
    avg_last = 'X' if len(subjects) >= (SUBJ_LAST_COL - SUBJ_FIRST_COL + 1) \
        else SUBJ_AVG_LAST
    for idx, student in enumerate(students):
        row = FIRST_DATA_ROW + idx
        for col in range(1, MAX_COL + 1):
            cell = ws.cell(row=row, column=col)
            cell._style = copy(data_styles[col])
        ws.row_dimensions[row].height = ROW_HEIGHT
        ws.cell(row=row, column=1, value=idx + 1)                           # A
        if row == FIRST_DATA_ROW:
            pass                                                            # B6 set in _fill_header
        else:
            ws.cell(row=row, column=2).value = (
                '=IF(Таблица1[[#This Row],[Столбец3]]<>"",'
                f'B{row - 1},"")'
            )
        ws.cell(row=row, column=3, value=student.full_name)                 # C
        if student.gender == 'F':
            ws.cell(row=row, column=4, value='д')                           # D
        elif student.gender == 'M':
            ws.cell(row=row, column=4, value='п')
        for sidx, subject in enumerate(subjects):
            grade = grades.get((student.id, subject))
            if grade is not None:
                ws.cell(row=row, column=SUBJ_FIRST_COL + sidx, value=grade)
        for col, formula in _row_formulas(row, avg_last).items():
            ws.cell(row=row, column=col).value = formula


def _min_chain(first, last, threshold):
    """'=IF(MIN(E6:X6)>=9,1,0)+IF(MIN(E7:X7)>=9,1,0)+...' — the owner's
    per-row chain counting students whose minimum grade reaches a band."""
    parts = [
        f'IF(MIN(E{r}:X{r})>={threshold},1,0)' for r in range(first, last + 1)
    ]
    return '=' + '+'.join(parts)


def _fill_footer(ws, spec, merges, delta, last, school):
    """Write the footer block shifted by `delta` rows: captured styles and
    label text verbatim, formulas regenerated against the live range."""
    first = FIRST_DATA_ROW
    s = FOOTER_FIRST_ROW + delta          # band rows s..s+3 (were 39..42)
    avg_row = s + 4                       # Холи миёна        (was 43)
    pct_row = s + 5                       # Фоиз              (was 44)
    cnt_row = s + 6                       # Шумораи хонандагон (was 45)
    sig1_row = FOOTER_FIRST_ROW + 8 + delta   # (was 47)
    sig2_row = FOOTER_FIRST_ROW + 10 + delta  # (was 49)

    for old_row, cells in spec.items():
        new_row = old_row + delta
        ws.row_dimensions[new_row].height = ROW_HEIGHT
        for col, (style, value) in cells.items():
            cell = ws.cell(row=new_row, column=col)
            cell._style = copy(style)
            cell.value = value

    for m in merges:
        ws.merge_cells(
            start_row=m.min_row + delta, start_column=m.min_col,
            end_row=m.max_row + delta, end_column=m.max_col,
        )

    # Band-count rows s..s+3, subject columns E..W (X carries no count
    # formulas in the template — the spare column is tallied in Z only).
    thresholds = ('>=9', '>=7', '>=4')
    # Subtraction order mirrors the owner verbatim: row s+1 subtracts
    # E{s}, row s+2 subtracts E{s+1} then E{s}.
    for i, threshold in enumerate(thresholds):
        row = s + i
        for col in range(SUBJ_FIRST_COL, SUBJ_LAST_COL):
            letter = get_column_letter(col)
            tail = ''.join(
                f'-{letter}{s + j}' for j in range(i - 1, -1, -1))
            ws.cell(row=row, column=col).value = (
                f'=COUNTIF({letter}${first}:{letter}{last},"{threshold}"){tail}'
            )
    for col in range(SUBJ_FIRST_COL, SUBJ_LAST_COL):
        letter = get_column_letter(col)
        ws.cell(row=s + 3, column=col).value = (
            f'=COUNTIF({letter}${first}:{letter}{last},"<4")'
        )

    # Z: students whose MINIMUM subject grade stays inside each band.
    ws.cell(row=s, column=26).value = _min_chain(first, last, 9)
    ws.cell(row=s + 1, column=26).value = _min_chain(first, last, 7) + f'-Z{s}'
    ws.cell(row=s + 2, column=26).value = (
        _min_chain(first, last, 4) + f'-Z{s}-Z{s + 1}'
    )
    ws.cell(row=s + 3, column=26).value = (
        f'=IF(COUNTIF(E{avg_row}:X{avg_row},">0"),'
        f'Z{cnt_row}-Z{s}-Z{s + 1}-Z{s + 2},"")'
    )

    # Холи миёна row.
    for col in range(SUBJ_FIRST_COL, SUBJ_LAST_COL):
        letter = get_column_letter(col)
        ws.cell(row=avg_row, column=col).value = (
            f'=IF(COUNTIF({letter}{s}:{letter}{s + 3},">0"),'
            f'AVERAGE({letter}{first}:{letter}{last}),"")'
        )
    ws.cell(row=avg_row, column=26).value = (
        f'=IF(COUNTIF(E{avg_row}:Y{avg_row},">0"),'
        f'AVERAGE(E{avg_row}:Y{avg_row}),"")'
    )
    ws.cell(row=avg_row, column=27).value = (
        f'=IF(COUNTIF(E{avg_row}:X{avg_row},">0"),'
        f'IF(AND(Z{avg_row}=AG{avg_row}, Z{pct_row}=AG{pct_row}),'
        f' "Ҳисоб дақиқ аст", "Хатои ҳисоббарорӣ"),"")'
    )
    ws.cell(row=avg_row, column=33).value = (
        f'=IF(COUNTIF(E{avg_row}:X{avg_row},">0"),'
        f'AVERAGE(AG{first}:AG{s + 3}),"")'
    )
    ws.cell(row=avg_row, column=34).value = (
        f'=IF(COUNTIF(AH{first}:AH{last},">0"),'
        f'AVERAGE(AH{first}:AH{last}),"")'
    )
    ws.cell(row=avg_row, column=35).value = (
        f'=IF(COUNTIF(AI{first}:AI{last},">0"),'
        f'AVERAGE(AI{first}:AI{last}),"")'
    )

    # Фоиз row (average * 10).
    for col in range(SUBJ_FIRST_COL, SUBJ_LAST_COL):
        letter = get_column_letter(col)
        ws.cell(row=pct_row, column=col).value = (
            f'=IF(COUNTIF({letter}{first}:{letter}{last},">0"),'
            f'{letter}{avg_row}*10,"")'
        )
    ws.cell(row=pct_row, column=26).value = (
        f'=IF(COUNTIF(E{pct_row}:Y{pct_row},">0"),'
        f'AVERAGE(E{pct_row}:Y{pct_row}),"")'
    )
    ws.cell(row=pct_row, column=33).value = (
        f'=IF(COUNTIF(E{avg_row}:X{avg_row},">0"),AG{avg_row}*10,"")'
    )

    # Шумораи хонандагон row + girls count.
    for col in range(SUBJ_FIRST_COL, SUBJ_LAST_COL):
        letter = get_column_letter(col)
        ws.cell(row=cnt_row, column=col).value = (
            f'=IF(COUNTIF({letter}{first}:{letter}{last},">0"),'
            f'{letter}{s}+{letter}{s + 1}+{letter}{s + 2},"")'
        )
    ws.cell(row=cnt_row, column=26).value = (
        f'=IF(COUNTIF(E{cnt_row}:W{cnt_row},">0"),'
        f'AVERAGE(E{cnt_row}:W{cnt_row}),"")'
    )
    ws.cell(row=cnt_row, column=32).value = (
        f'=IF(COUNTIF(D{first}:D{last},"<>"""),'
        f'COUNTIF(D{first}:D{last},"=д"),"")'
    )

    # Signature lines. The deputy-director name has no reliable portal
    # source and stays blank; the director comes from School.director.
    ws.cell(row=sig1_row, column=1).value = (
        'Муовини директор оид бо корҳои таълим:    _______________________'
    )
    director = _director_name(school)
    ws.cell(row=sig2_row, column=1).value = (
        f'                               Директори {_school_short_name(school)}:'
        f'    _______________________  {director}'.rstrip()
    )


def _resize_table(ws, last):
    """Point Таблица1 at the actual data range so its structured
    references ([Столбец3] etc.) keep working."""
    tbl = ws.tables.get('Таблица1')
    if tbl is None:
        return
    tbl.ref = f'A5:AJ{last}'
    if tbl.autoFilter is not None:
        tbl.autoFilter.ref = tbl.ref


def _page_setup(ws, last_row):
    """One A4 landscape page: template orientation/margins kept,
    fit-to-page and print area set for immediate printing."""
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_area = f'A1:AJ{last_row}'
    ws.print_options.horizontalCentered = True
