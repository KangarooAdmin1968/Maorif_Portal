"""Ҳуҷҷатнигорӣ → Ҳисоботи маориф — А.Ҷ.Ч. district summary Excel export.

Reproduces the education department's paper table: one row per academic
school, the per-student quarter-average band distribution (10-9 / 8-7 /
6-4 / 3-1), the сифатнокӣ share and the curriculum-mastery (4-10) share,
plus a Хамагӣ total row, the level-1 Эзоҳ note and signature lines.

collect_maorif_stats() returns plain JSON-serializable data on purpose —
when the future «чорак якунланиш» closure lands, that payload can be
frozen as a snapshot so a finalized quarter's report can never change.
"""
import io
import re
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from .class_report import _school_short_name
from .models import QuarterGrade, School, Student
from .utils import academic_schools, class_numeric_part, get_school_number
from .zavuch_report import _band

LEVELS = range(2, 12)          # the table counts grade levels 2..11
QUARTER_WORD_ORDINALS = {1: 'якуми', 2: 'дуюми', 3: 'сеюми', 4: 'чоруми'}

FIRST_DATA_ROW = 6
LAST_COL = 14                  # A..N

_thin = Side(style='thin')
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)


def _district_sort_key(school):
    """МТМУ №1..№N first, then Гимназия, then Литсей — the paper order."""
    slug = get_school_number(school)
    m = re.fullmatch(r'([a-z]*)(\d+)', slug)
    prefix = m.group(1) if m else 'x'
    num = int(m.group(2)) if m else 9999
    order = {'': 0, 'g': 1, 'l': 2}.get(prefix, 3)
    return (order, num, school.name or '')


def collect_maorif_stats(quarter, today=None):
    """District stats for one quarter — JSON-friendly plain data.

    Built so the future quarter-closure button can persist this payload
    verbatim; the workbook builder below only reads it.
    """
    from .views import academic_year_start  # local import avoids cycles

    ay = academic_year_start(today)
    schools = sorted(academic_schools(), key=_district_sort_key)
    school_ids = [s.id for s in schools]

    # Roster: levels 2-11 feed the table; level 1 feeds the Эзоҳ note.
    where = {}                       # student_id -> (school_id, class_name)
    roster = {s.id: {'total': 0} for s in schools}
    level1_total = level1_girls = 0
    for row in Student.objects.filter(school_id__in=school_ids).values(
            'id', 'school_id', 'class_name', 'gender'):
        lvl = class_numeric_part(row['class_name'])
        if lvl == 1:
            level1_total += 1
            if row['gender'] == 'F':
                level1_girls += 1
        if lvl in LEVELS:
            roster[row['school_id']]['total'] += 1
            where[row['id']] = (row['school_id'], row['class_name'])

    # Per-student quarter average -> band, attributed to the school.
    sums = {}
    for row in QuarterGrade.objects.filter(
            student__school_id__in=school_ids, quarter=quarter,
            academic_year=ay, grade__isnull=False).values(
            'student_id', 'grade'):
        entry = sums.setdefault(row['student_id'], [0.0, 0])
        entry[0] += float(row['grade'])
        entry[1] += 1

    stats = {s.id: {'graded': 0, 'hi': 0, 'md': 0, 'rm': 0, 'lo': 0}
             for s in schools}
    for sid, (total, count) in sums.items():
        loc = where.get(sid)
        if not loc:
            continue
        st = stats[loc[0]]
        st['graded'] += 1
        st[_band(total / count)] += 1

    # Ungraded 2-11 students per school+class — the Эзоҳ footnote lists
    # them the same way the paper form does.
    ungraded = defaultdict(int)
    for sid, (school_id, class_name) in where.items():
        if sid not in sums:
            ungraded[(school_id, class_name)] += 1

    by_id = {s.id: s for s in schools}
    return {
        'quarter': quarter,
        'academic_year': ay,
        'level1_total': level1_total,
        'level1_girls': level1_girls,
        'rows': [{
            'school_id': s.id,
            'school_name': _school_short_name(s),
            'total': roster[s.id]['total'],
            **stats[s.id],
        } for s in schools],
        'ungraded': sorted(
            ({'school_name': _school_short_name(by_id[sid]),
              'class_name': cn, 'n': n}
             for (sid, cn), n in ungraded.items()),
            key=lambda r: (r['school_name'], class_numeric_part(
                r['class_name']), r['class_name'])),
    }


def _district_director_name():
    dept = School.objects.filter(type='Идораи маориф').first()
    return (dept.director or '').strip() if dept else ''


def _pct(part, whole):
    """Rounded part/whole*100 as a plain value — displays even in
    Excel's Protected View where formulas are never recalculated."""
    return round(part / whole * 100, 2) if whole else 0


def _write_title(ws, quarter, ay):
    ordinal = QUARTER_WORD_ORDINALS.get(quarter, f'{quarter}-ум')
    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1, end_column=LAST_COL)
    ws['A1'] = ('дар бораи натиҷаи арзёбии ҷамъбастии чоряки '
                f'{ordinal} (А.Ҷ.Ч-{quarter})')
    ws['A1'].font = Font(italic=True, size=11)
    ws['A1'].alignment = LEFT
    ws.merge_cells(start_row=2, start_column=1,
                   end_row=2, end_column=LAST_COL)
    ws['A2'] = 'МАЪЛУМОТ'
    ws['A2'].font = Font(bold=True, size=14)
    ws['A2'].alignment = CENTER
    ws.merge_cells(start_row=3, start_column=1,
                   end_row=3, end_column=LAST_COL)
    ws['A3'] = ('хонандагони муассисахои таъсилоти умумӣ дар соли '
                f'таҳсили {ay}-{ay + 1} оид ба сатҳу сифати дониши '
                'хонандагон дар ноҳияи Зафаробод')
    ws['A3'].font = Font(bold=True, size=11)
    ws['A3'].alignment = CENTER


def _write_header(ws):
    """Rows 4-5: merged captions + band sub-headers."""

    def merge_header(text, col, last_col=None):
        ws.merge_cells(start_row=4, start_column=col, end_row=5,
                       end_column=last_col or col)
        c = ws.cell(row=4, column=col, value=text)
        c.font = Font(bold=True, size=10)
        c.alignment = CENTER

    merge_header('№', 1)
    merge_header('Ноҳияи Зафаробод', 2)
    merge_header('Шумораи хонандагони синфи 2-11', 3)
    merge_header('Иштирок намуд', 4)
    # «Баҳо» is a single-row group label — row 5 keeps the band sub-heads.
    ws.merge_cells(start_row=4, start_column=5, end_row=4, end_column=12)
    c = ws.cell(row=4, column=5, value='Баҳо')
    c.font = Font(bold=True, size=10)
    c.alignment = CENTER
    merge_header('сифатнокӣ', 13)
    merge_header('Аҳ. худкунии барномаҳои таълимӣ (4-10)', 14)
    for col, text in zip(range(5, 13),
                         ('10-9', '%', '8-7', '%', '6-4', '%', '3-1', '%')):
        c = ws.cell(row=5, column=col, value=text)
        c.font = Font(bold=True, size=10)
        c.alignment = CENTER
    # Merged cells never auto-grow — reserve height so the wrapped
    # captions (esp. the long col-N one) stay fully visible.
    ws.row_dimensions[4].height = 42
    ws.row_dimensions[5].height = 15


def _write_data_row(ws, r, index, name, data, bold=None):
    """One table row: counts + computed percents for `data` keys
    (total/graded/hi/md/rm/lo). index=None renders an unnumbered bold
    row (Хамагӣ); the totals-only sheet passes index=1 + bold=True."""
    if bold is None:
        bold = index is None
    if index is not None:
        ws.cell(row=r, column=1, value=index)
    name_cell = ws.cell(row=r, column=2, value=name)
    name_cell.alignment = LEFT
    if bold:
        name_cell.font = Font(bold=True)
    ws.cell(row=r, column=3, value=data['total'])
    ws.cell(row=r, column=4, value=data['graded'])
    for col, key in ((5, 'hi'), (7, 'md'), (9, 'rm'), (11, 'lo')):
        c = ws.cell(row=r, column=col, value=data[key])
        if bold:
            c.font = Font(bold=True)
    g = data['graded']
    for col, part in ((6, data['hi']), (8, data['md']),
                      (10, data['rm']), (12, data['lo'])):
        c = ws.cell(row=r, column=col, value=_pct(part, g))
        c.number_format = '0.00'
        if bold:
            c.font = Font(bold=True)
    for col, value in (
            (13, _pct(data['hi'] + data['md'], g)),
            # 4-10 mastery = participants who never landed in the 1-3 band.
            (14, _pct(g - data['lo'], g))):
        c = ws.cell(row=r, column=col, value=value)
        c.number_format = '0.00'
        if bold:
            c.font = Font(bold=True)


def _border_table(ws, last_row):
    for r in range(4, last_row + 1):
        for col in range(1, LAST_COL + 1):
            cell = ws.cell(row=r, column=col)
            cell.border = BORDER
            if cell.alignment.horizontal is None:
                cell.alignment = CENTER


def _write_footer(ws, start_row, stats, prepared_by):
    ordinal = QUARTER_WORD_ORDINALS.get(
        stats['quarter'], f"{stats['quarter']}-ум")
    row = start_row

    def line(text, font):
        nonlocal row
        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=LAST_COL)
        c = ws.cell(row=row, column=1, value=text)
        c.alignment = LEFT
        c.font = font
        row += 1

    line('Эзоҳ: Шумораи хонандагони синфи 1-ум '
         f"{stats['level1_total']} нафар, аз он ҷумла духтар "
         f"{stats['level1_girls']} нафарро ташкил медиҳад.",
         Font(size=10, italic=True))

    if stats['ungraded']:
        # One row per school, classes compact — errors stay scannable
        # without one giant wrapped sentence.
        line(f'Баҳо нагирифтаанд дар чоряки {ordinal}:',
             Font(size=10, italic=True, bold=True))
        by_school = defaultdict(list)
        for u in stats['ungraded']:
            by_school[u['school_name']].append(u)
        order = [r['school_name'] for r in stats['rows']
                 if r['school_name'] in by_school]
        order += [n for n in by_school if n not in set(order)]
        for school_name in order:
            items = by_school[school_name]
            listed = ', '.join(f"{u['class_name']} ({u['n']})"
                               for u in items)
            line(f'{school_name}: {listed} — жами '
                 f"{sum(u['n'] for u in items)} нафар",
                 Font(size=9, italic=True))

    row += 1
    ws.merge_cells(start_row=row, start_column=1,
                   end_row=row, end_column=6)
    ws.cell(row=row, column=1,
            value=f'Мудири шуъбаи маориф:    {_district_director_name()}'
                  .rstrip())
    ws.merge_cells(start_row=row + 1, start_column=1,
                   end_row=row + 1, end_column=6)
    ws.cell(row=row + 1, column=1,
            value=f'Ичрокунанда:    {prepared_by}'.rstrip())


def _layout(ws):
    widths = {1: 5, 2: 26, 3: 14, 4: 12, 13: 11, 14: 18}
    for col in range(1, LAST_COL + 1):
        ws.column_dimensions[get_column_letter(col)].width = (
            widths.get(col, 8))
    # Everything fits one A4 landscape sheet — scale to a single page.
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5


def _finish(wb):
    wb.calculation.fullCalcOnLoad = True
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _totals(stats):
    return {k: sum(r[k] for r in stats['rows'])
            for k in ('total', 'graded', 'hi', 'md', 'rm', 'lo')}


def build_maorif_report_workbook(quarter, prepared_by=''):
    """District А.Ҷ.Ч. workbook, one row per academic school -> BytesIO."""
    stats = collect_maorif_stats(quarter)
    wb = Workbook()
    ws = wb.active
    ws.title = f'АҶЧ-{quarter}'

    _write_title(ws, quarter, stats['academic_year'])
    _write_header(ws)
    for i, row in enumerate(stats['rows']):
        _write_data_row(ws, FIRST_DATA_ROW + i, i + 1,
                        row['school_name'], row)
    tot = FIRST_DATA_ROW + len(stats['rows'])
    _write_data_row(ws, tot, None, 'Хамагӣ', _totals(stats))
    _border_table(ws, tot)
    _write_footer(ws, tot + 2, stats, prepared_by)
    _layout(ws)
    return _finish(wb)


def build_maorif_total_workbook(quarter, prepared_by=''):
    """Same district summary collapsed to the single Хамагӣ row.

    This is the standalone totals page the department sends upward —
    identical columns and footer to the per-school workbook.
    """
    stats = collect_maorif_stats(quarter)
    wb = Workbook()
    ws = wb.active
    ws.title = f'АҶЧ-{quarter} — ҳамагӣ'

    _write_title(ws, quarter, stats['academic_year'])
    _write_header(ws)
    _write_data_row(ws, FIRST_DATA_ROW, 1, 'Хамагӣ', _totals(stats),
                    bold=True)
    _border_table(ws, FIRST_DATA_ROW)
    _write_footer(ws, FIRST_DATA_ROW + 2, stats, prepared_by)
    _layout(ws)
    return _finish(wb)
