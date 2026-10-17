"""Ҳуҷҷатнигорӣ → Ҳисоботи синфҳо — quarterly class report export.

Covers: the Documents-area tab and class picker, school isolation and the
табел audience gate, student/quarter-grade mapping onto the owner's
template cells, formula/layout preservation, Таблица1 resizing, and the
one-page A4 landscape print setup.
"""
import datetime
import io
import re

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from portal.class_report import (
    TEMPLATE_PATH, build_class_report_workbook, report_subjects,
)
from portal.curriculum_hours import UNGRADED_SUBJECTS
from portal.models import normalize_class_name
from portal.tests.helpers import (
    MATH, TAJIK, GoldenBase, make_class_subject, make_quarter_grade,
    make_student, make_teacher_profile, make_user,
)

TARBIYAVI = 'СОАТИ ТАРБИЯВӢ'


def _subject_col(ws, subject):
    """Find the report column whose row-3 header is `subject` (E..X)."""
    for col in range(5, 25):
        if ws.cell(row=3, column=col).value == subject:
            return col
    return None


class ClassReportBuilderTests(GoldenBase):
    """portal.class_report.build_class_report_workbook."""

    def _build(self, school=None, class_name='7-А', quarter=1):
        return load_workbook(build_class_report_workbook(
            school or self.school_a, class_name, quarter,
            today=datetime.date(2025, 10, 15)))

    def setUp(self):
        super().setUp()
        make_quarter_grade(self.s1, '7-А', MATH, quarter=1, grade=9)
        make_quarter_grade(self.s1, '7-А', TAJIK, quarter=1, grade=7)
        make_quarter_grade(self.s1, '7-А', MATH, quarter=2, grade=4)
        make_quarter_grade(self.s2, '7-А', MATH, quarter=1, grade=5)
        make_quarter_grade(self.s2, '7-А', TAJIK, quarter=1, grade=10)
        # School B data must never leak into a school-A report.
        make_quarter_grade(self.s3, '7-Б', MATH, quarter=1, grade=10)

    # -- selection / data -------------------------------------------------

    def test_selected_class_and_students_only(self):
        ws = self._build().worksheets[0]
        names = [ws.cell(row=r, column=3).value for r in range(6, 8)]
        self.assertEqual(sorted(names), ['Мадина Каримова', 'Фирӯз Раҳимов'])
        self.assertNotIn('Ҷаҳонгир Назаров', names)
        self.assertEqual(ws['B6'].value, '7-А')
        self.assertEqual(ws.title, '7-А')

    def test_quarter_selects_matching_grades(self):
        ws1 = self._build(quarter=1).worksheets[0]
        ws2 = self._build(quarter=2).worksheets[0]
        math = _subject_col(ws1, MATH)
        self.assertIsNotNone(math)
        letter = get_column_letter(math)
        row_s1 = 6 if ws1['C6'].value == 'Фирӯз Раҳимов' else 7
        row_s2 = 6 if ws1['C6'].value == 'Мадина Каримова' else 7
        self.assertEqual(ws1[f'{letter}{row_s1}'].value, 9)
        self.assertEqual(ws1[f'{letter}{row_s2}'].value, 5)
        self.assertEqual(ws2[f'{letter}{row_s1}'].value, 4)
        self.assertIsNone(ws2[f'{letter}{row_s2}'].value)

    def test_gender_and_row_numbers(self):
        ws = self._build().worksheets[0]
        gender = {ws.cell(row=r, column=3).value: ws.cell(row=r, column=4).value
                  for r in (6, 7)}
        self.assertEqual(gender['Мадина Каримова'], 'д')
        self.assertEqual(gender['Фирӯз Раҳимов'], 'п')
        self.assertEqual([ws.cell(row=r, column=1).value for r in (6, 7)],
                         [1, 2])

    def test_subject_headers_are_class_subjects(self):
        ws = self._build().worksheets[0]
        headers = {ws.cell(row=3, column=c).value for c in range(5, 25)}
        headers.discard(None)
        self.assertEqual(headers, set(report_subjects(self.school_a, '7-А')))
        self.assertIn(MATH, headers)
        # СОАТИ ТАРБИЯВӢ is an ungraded subject — never a report column.
        self.assertFalse(set(UNGRADED_SUBJECTS) & headers)

    def test_title_and_school_line(self):
        ws = self._build().worksheets[0]
        self.assertIn('чоряки 1-ум дар синфи', ws['A2'].value)
        self.assertIn('мактаб №1', ws['A2'].value.lower())
        self.assertIn('2025-2026', ws['A2'].value)
        # Template A3's INDIRECT evaluates to #REF! — replaced by the row
        # header label instead of copying a broken formula.
        self.assertEqual(ws['A3'].value, '№')

    def test_class_teacher_from_tarbiyavi_allocation(self):
        profile = make_teacher_profile(self.teacher_a, self.school_a,
                                       full_name='Раҳимова О.')
        make_class_subject(self.school_a, '7-А', TARBIYAVI,
                           allocated=profile)
        ws = self._build().worksheets[0]
        self.assertEqual(ws['AG1'].value, 'Раҳимова О.')

    def test_signature_uses_school_director(self):
        self.school_a.director = 'Одинаев С.З.'
        self.school_a.save()
        ws = self._build().worksheets[0]
        sig = ws['A18'].value  # 49 shifted by delta = (6+2)-39 = -31+... = 18
        self.assertIn('Одинаев С.З.', sig)
        self.assertIn('Директори', sig)

    # -- formulas ---------------------------------------------------------

    def test_student_band_formulas_preserved(self):
        ws = self._build().worksheets[0]
        self.assertEqual(
            ws['Y6'].value,
            '=IF(COUNT($E6:$X6)>0,COUNTIF($E6:$X6,">=9"),"")')
        self.assertEqual(
            ws['Z6'].value,
            '=IF(COUNT($E6:$X6)>0,Y6/COUNT($E6:$X6)*100,"")')
        self.assertEqual(
            ws['AA6'].value,
            '=IF(COUNT($E6:$X6)>0,COUNTIF($E6:$X6,">=7")-Y6,"")')
        self.assertEqual(
            ws['AG6'].value, '=IF(COUNT(E6:W6)>0,AVERAGE(E6:W6),"")')
        self.assertEqual(ws['AH6'].value, '=IF(COUNT(E6:W6)>0,100,"")')
        self.assertEqual(ws['AI6'].value, '=IF(COUNT(E6:W6)>0,Z6+AB6,"")')

    def test_footer_formulas_scaled_to_class_size(self):
        # 2 students -> data rows 6..7, band block starts at row 8.
        ws = self._build().worksheets[0]
        math = get_column_letter(_subject_col(ws, MATH))
        self.assertEqual(ws[f'{math}8'].value,
                         f'=COUNTIF({math}$6:{math}7,">=9")')
        self.assertEqual(ws[f'{math}9'].value,
                         f'=COUNTIF({math}$6:{math}7,">=7")-{math}8')
        # Owner order: the 4-6 row subtracts s+1 first, then s.
        self.assertEqual(ws[f'{math}10'].value,
                         f'=COUNTIF({math}$6:{math}7,">=4")-{math}9-{math}8')
        self.assertEqual(ws[f'{math}11'].value,
                         f'=COUNTIF({math}$6:{math}7,"<4")')
        self.assertEqual(
            ws['Z8'].value,
            '=IF(MIN(E6:X6)>=9,1,0)+IF(MIN(E7:X7)>=9,1,0)')
        self.assertEqual(ws['A8'].value, 'Хол')
        self.assertEqual(ws['C8'].value, '9-10')
        self.assertEqual(ws['A12'].value, 'Холи миёна')
        self.assertEqual(ws['A13'].value, 'Фоиз')
        self.assertEqual(ws['A14'].value, 'Шумораи хонандагон')
        self.assertIn('Муовини директор', ws['A16'].value)
        self.assertIn('Директори', ws['A18'].value)

    def test_band_counts_match_source_data(self):
        """Independent check against source data: quarter grades land in
        the cell at (student row, subject column); subjects without grades
        stay blank so the owner's COUNTIF/IF guards render "" not 0."""
        ws = self._build().worksheets[0]
        math, tajik = _subject_col(ws, MATH), _subject_col(ws, TAJIK)
        cols = {math: [9, 5], tajik: [7, 10]}
        for col, grades in cols.items():
            cells = [ws.cell(row=r, column=col).value for r in (6, 7)]
            self.assertEqual(sorted(cells), sorted(grades))
        graded = {math, tajik}
        for col in range(5, 25):
            if ws.cell(row=3, column=col).value and col not in graded:
                self.assertIsNone(ws.cell(row=6, column=col).value)

    # -- layout preservation ----------------------------------------------

    def test_layout_preserved(self):
        ws = self._build().worksheets[0]
        template = load_workbook(TEMPLATE_PATH).worksheets[0]
        for col in ('A', 'C', 'E', 'Y', 'AJ'):
            self.assertEqual(ws.column_dimensions[col].width,
                             template.column_dimensions[col].width)
        self.assertEqual(ws.row_dimensions[4].height, 36)
        self.assertTrue(ws.row_dimensions[5].hidden)
        self.assertEqual(ws['E3'].alignment.textRotation, 90)
        # Footer merges shift with class size: 2 students -> band merge
        # A8:B11 (was A39:B42), signature merges A16:AJ16 + A18:AJ18.
        merged = {str(m) for m in ws.merged_cells.ranges}
        self.assertIn('A8:B11', merged)
        self.assertIn('A16:AJ16', merged)
        self.assertIn('A18:AJ18', merged)
        self.assertEqual(ws.tables['Таблица1'].ref, 'A5:AJ7')

    def test_bigger_class_shifts_footer_down(self):
        make_student(self.school_a, '7-А', 'Ғаффоров Аҳмад')
        ws = self._build().worksheets[0]
        # 3 students -> last data row 8, band block 9..12.
        self.assertEqual(ws['A9'].value, 'Хол')
        merged = {str(m) for m in ws.merged_cells.ranges}
        self.assertIn('A9:B12', merged)
        self.assertIn('A17:AJ17', merged)
        self.assertEqual(ws.tables['Таблица1'].ref, 'A5:AJ8')

    # -- print setup -------------------------------------------------------

    def test_single_page_a4_landscape_print(self):
        ws = self._build().worksheets[0]
        self.assertEqual(ws.page_setup.orientation, 'landscape')
        self.assertEqual(int(ws.page_setup.paperSize), 9)  # A4
        self.assertEqual(ws.page_setup.fitToWidth, 1)
        self.assertEqual(ws.page_setup.fitToHeight, 1)
        self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage)
        self.assertTrue(ws.print_options.horizontalCentered)
        # Print area covers the whole report through the signature row.
        self.assertTrue(str(ws.print_area).endswith('$A$1:$AJ$18'))


class ClassReportPageTests(GoldenBase):
    """GET /schools/documents/class-reports/ + export endpoint."""

    def setUp(self):
        super().setUp()
        self.url = reverse('class_reports')
        self.export_url = reverse('class_report_export')

    def test_zavuch_lists_own_school_classes(self):
        make_student(self.school_a, '1-А', 'Хуршед Одилов')
        self.client.force_login(self.zavuch_a)
        r = self.client.get(self.url)
        self.assertEqual(r.status_code, 200)
        listed = [c['class_name'] for c in r.context['classes']]
        self.assertEqual(listed, ['7-А'])  # non-graded 1-А excluded

    def test_regular_teacher_denied(self):
        self.client.force_login(self.teacher_a)
        for url in (self.url,
                    f'{self.export_url}?class_name=7-А'):
            r = self.client.get(url)
            self.assertEqual(r.status_code, 302)
            self.assertIn(reverse('dashboard'), r.url)

    def test_anonymous_redirected_to_login(self):
        r = self.client.get(self.url)
        self.assertEqual(r.status_code, 302)
        self.assertIn('login', r.url)

    def test_superuser_picks_school(self):
        self.client.force_login(self.admin)
        r = self.client.get(self.url, {'school_id': self.school_b.id})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context['school'], self.school_b)
        self.assertEqual([c['class_name'] for c in r.context['classes']],
                         ['7-Б'])

    def test_documents_tab_links_to_class_reports(self):
        self.client.force_login(self.zavuch_a)
        r = self.client.get(reverse('school_documents'))
        self.assertContains(r, 'Ҳисоботи синфҳо')
        self.assertContains(r, reverse('class_reports'))

    def test_export_returns_workbook_for_own_school(self):
        make_quarter_grade(self.s1, '7-А', MATH, quarter=1, grade=9)
        self.client.force_login(self.zavuch_a)
        r = self.client.get(
            self.export_url, {'class_name': '7-А', 'quarter': 1})
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.content)).worksheets[0]
        self.assertEqual(ws.title, '7-А')
        math = get_column_letter(_subject_col(ws, MATH))
        row = 6 if ws['C6'].value == 'Фирӯз Раҳимов' else 7
        self.assertEqual(ws[f'{math}{row}'].value, 9)

    def test_export_school_isolation(self):
        """zavuch_1 (school A) cannot pull school B's 7-Б even by
        forging school_id — the school resolves to their own."""
        self.client.force_login(self.zavuch_a)
        r = self.client.get(self.export_url, {
            'school_id': self.school_b.id, 'class_name': '7-Б',
            'quarter': 1})
        self.assertEqual(r.status_code, 302)
        self.assertIn(reverse('class_reports'), r.url)
        # And school B's zavuch gets only school-B data.
        zavuch_b = make_user('zavuch_2', role='teacher',
                             school=self.school_b)
        self.client.force_login(zavuch_b)
        r = self.client.get(self.export_url, {'class_name': '7-Б',
                                              'quarter': 1})
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.content)).worksheets[0]
        self.assertEqual(ws['C6'].value, 'Ҷаҳонгир Назаров')

    def test_export_rejects_unknown_and_nongraded_class(self):
        make_student(self.school_a, '1-А', 'Хуршед Одилов')
        self.client.force_login(self.zavuch_a)
        for cls in ('9-З', '1-А'):
            r = self.client.get(self.export_url,
                                {'class_name': cls, 'quarter': 1})
            self.assertEqual(r.status_code, 302)
            self.assertIn(reverse('class_reports'), r.url)

    def test_export_content_disposition(self):
        self.client.force_login(self.zavuch_a)
        r = self.client.get(self.export_url,
                            {'class_name': '7-А', 'quarter': 1})
        self.assertIn('attachment', r['Content-Disposition'])
        self.assertIn('.xlsx', r['Content-Disposition'])
