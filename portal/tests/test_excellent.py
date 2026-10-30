"""⭐ Аълохон — «Рӯйхати хонандагони аълохон» document + Excel export.

A student makes the list only when EVERY graded subject of their class
carries a quarter mark of 9 or 10 — a missing mark disqualifies too.
"""
import io

from django.urls import reverse
from openpyxl import load_workbook

from portal.curriculum_hours import UNGRADED_SUBJECTS
from portal.excellent import collect_excellent_students
from portal.models import ClassSubject, QuarterGrade, normalize_subject

from .helpers import (D_Q1, GoldenBase, make_grade, make_quarter_grade,
                      make_student, make_user)


class ExcellentStudentsTests(GoldenBase):

    def _all_graded(self, school, class_name):
        """Active graded ClassSubjects of a class — the list's demand."""
        return {
            normalize_subject(cs.subject)
            for cs in ClassSubject.objects.filter(
                school=school, class_name=class_name, is_active=True)
            if normalize_subject(cs.subject) not in UNGRADED_SUBJECTS
        }

    def _fill(self, student, school, quarter=1, mark=9, skip=None):
        """Quarter marks in every graded subject (skip one if asked)."""
        for subj in self._all_graded(school, student.class_name):
            if skip and subj == normalize_subject(skip):
                continue
            make_quarter_grade(student, student.class_name, subj,
                               quarter=quarter, grade=mark)

    def test_all_marks_nine_and_ten_make_the_list(self):
        self._fill(self.s1, self.school_a)
        # A mix of 9s and 10s is fine — bump one subject to 10.
        make_quarter_grade(self.s1, '7-А', 'МАТЕМАТИКА',
                           quarter=1, grade=10)
        data = collect_excellent_students(self.school_a, 1)
        self.assertIn(self.s1, data['students'])
        self.assertNotIn(self.s2, data['students'])   # no marks at all

    def test_one_low_mark_excludes(self):
        self._fill(self.s1, self.school_a)
        make_quarter_grade(self.s1, '7-А', 'МАТЕМАТИКА',
                           quarter=1, grade=8)
        data = collect_excellent_students(self.school_a, 1)
        self.assertNotIn(self.s1, data['students'])

    def test_live_provisional_mark_fills_missing_row(self):
        # Mirrors the Ҳисобот: no stored QuarterGrade for МАТЕМАТИКА, but
        # a daily 10 makes the journal's Ҷ provisional value a 10 — the
        # student still makes the list.
        self._fill(self.s1, self.school_a, skip='МАТЕМАТИКА')
        make_grade(self.s1, subject='МАТЕМАТИКА', score=10, date=D_Q1)
        # Simulate the stored row never having been written (bulk import,
        # a failed recalc — however the gap arose in production).
        QuarterGrade.objects.filter(
            student=self.s1, subject='МАТЕМАТИКА', quarter=1).delete()
        data = collect_excellent_students(self.school_a, 1)
        self.assertIn(self.s1, data['students'])

    def test_stray_mark_outside_curriculum_does_not_exclude(self):
        # A stored mark in a subject her class does not study (an old
        # or deactivated ClassSubject) never shows in her Ҳисобот row —
        # so it must not keep her off the аъло list either.
        self._fill(self.s1, self.school_a)
        make_quarter_grade(self.s1, '7-А', 'ФАНИ НОМАЪЛУМ',
                           quarter=1, grade=4)
        data = collect_excellent_students(self.school_a, 1)
        self.assertIn(self.s1, data['students'])

    def test_missing_subject_excludes(self):
        self._fill(self.s1, self.school_a, skip='МАТЕМАТИКА')
        data = collect_excellent_students(self.school_a, 1)
        self.assertNotIn(self.s1, data['students'])

    def test_other_quarter_is_empty(self):
        self._fill(self.s1, self.school_a, quarter=1)
        data = collect_excellent_students(self.school_a, 2)
        self.assertEqual(data['students'], [])

    def test_grade_one_class_never_listed(self):
        s1a = make_student(self.school_a, '1-А', 'Хуршед Одилов')
        self._fill(s1a, self.school_a)
        data = collect_excellent_students(self.school_a, 1)
        self.assertNotIn(s1a, data['students'])

    def test_school_b_marks_do_not_leak(self):
        self._fill(self.s3, self.school_b, quarter=1)
        data = collect_excellent_students(self.school_a, 1)
        self.assertEqual(data['students'], [])


class ExcellentStudentsViewTests(GoldenBase):

    def _seed(self):
        for cs in ClassSubject.objects.filter(
                school=self.school_a, class_name='7-А', is_active=True):
            if normalize_subject(cs.subject) in UNGRADED_SUBJECTS:
                continue
            make_quarter_grade(self.s1, '7-А', cs.subject,
                               quarter=1, grade=9)

    def test_zavuch_sees_list_with_official_title(self):
        self._seed()
        self.client.force_login(self.zavuch_a)
        resp = self.client.get(reverse('excellent_students'), {'quarter': 1})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Рӯйхати хонандагони аълохон')
        self.assertContains(resp, self.school_a.name)
        self.assertContains(resp, 'Фирӯз Раҳимов')
        self.assertNotContains(resp, 'Мадина Каримова')

    def test_teacher_denied(self):
        self.client.force_login(self.teacher_a)
        resp = self.client.get(reverse('excellent_students'), {'quarter': 1})
        self.assertEqual(resp.status_code, 302)
        resp = self.client.get(
            reverse('excellent_students_export'), {'quarter': 1})
        self.assertEqual(resp.status_code, 302)

    def test_principal_can_view_own_school(self):
        self._seed()
        self.client.force_login(self.principal_a)
        resp = self.client.get(reverse('excellent_students'), {'quarter': 1})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Фирӯз Раҳимов')

    def test_maorif_staff_picks_school(self):
        # The маориф staff account (no school of its own) gets the same
        # picker as the director and can view/export any school.
        self._seed()
        staff = make_user('maorif_staff_ex', staff=True)
        self.client.force_login(staff)
        resp = self.client.get(reverse('excellent_students'), {'quarter': 1})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'name="school_id"')
        resp = self.client.get(reverse('excellent_students'), {
            'quarter': 1, 'school_id': self.school_a.id})
        self.assertContains(resp, 'Фирӯз Раҳимов')
        resp = self.client.get(reverse('excellent_students_export'), {
            'quarter': 1, 'school_id': self.school_a.id})
        self.assertEqual(resp.status_code, 200)

    def test_director_picks_school(self):
        self._seed()
        self.client.force_login(self.director)
        # No school chosen -> picker prompt, no list.
        resp = self.client.get(reverse('excellent_students'), {'quarter': 1})
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, 'Фирӯз Раҳимов')
        # ?school_id= scopes the director to one school.
        resp = self.client.get(reverse('excellent_students'), {
            'quarter': 1, 'school_id': self.school_a.id})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Фирӯз Раҳимов')
        # Export honors the same school param.
        resp = self.client.get(reverse('excellent_students_export'), {
            'quarter': 1, 'school_id': self.school_a.id})
        self.assertEqual(resp.status_code, 200)

    def test_export_workbook(self):
        self._seed()
        self.client.force_login(self.zavuch_a)
        resp = self.client.get(
            reverse('excellent_students_export'), {'quarter': 1})
        self.assertEqual(resp.status_code, 200)
        self.assertIn('attachment', resp['Content-Disposition'])
        ws = load_workbook(io.BytesIO(resp.content)).active
        self.assertIn('Рӯйхати хонандагони аълохон', ws['A1'].value)
        self.assertIn(self.school_a.name, ws['A1'].value)
        self.assertEqual(
            [ws.cell(row=4, column=c).value for c in range(1, 4)],
            ['№', 'Ному насаб', 'Синф'])
        self.assertEqual(
            [ws.cell(row=5, column=c).value for c in range(1, 4)],
            [1, 'Фирӯз Раҳимов', '7-А'])

    def test_export_empty_quarter_says_so(self):
        self.client.force_login(self.zavuch_a)
        resp = self.client.get(
            reverse('excellent_students_export'), {'quarter': 2})
        ws = load_workbook(io.BytesIO(resp.content)).active
        self.assertIn('аълохон ёфт нашуд', ws['A5'].value)
