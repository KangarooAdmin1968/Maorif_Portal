"""Phase 3 tests — teacher-scoped group access for journals, grades,
lessons and assessments.

Fixture: School A, class 7-А, subject МАТЕМАТИКА with two TeachingGroups:
    Гурӯҳи 1 -> teacher_a (also legacy cs.teacher) -> s1
    Гурӯҳи 2 -> teacher_a2 (group-only, no CS-level assignment) -> s2
"""
import io
import json

from django.urls import reverse
from openpyxl import load_workbook

from portal.models import (
    Assessment, Grade, Lesson, SubjectGroupMembership, TeachingGroup,
)
from portal.tests.helpers import (
    D_Q1, MATH, TAJIK, GoldenBase, make_class_subject, make_grade,
    make_student, make_teacher_profile, make_user,
)


class GroupedFixture(GoldenBase):
    """School A 7-А МАТЕМАТИКА grouped into two TeachingGroups."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.cs_math = cls.cs_a  # 7-А МАТЕМАТИКА, teacher=teacher_a
        cls.profile_a = make_teacher_profile(
            cls.teacher_a, cls.school_a, full_name='Омӯзгор А')
        cls.teacher_a2 = make_user(
            'teacher_3', role='teacher', school=cls.school_a)
        cls.profile_a2 = make_teacher_profile(
            cls.teacher_a2, cls.school_a, full_name='Омӯзгор Б')
        cls.s_extra = make_student(cls.school_a, '7-А', 'Наврӯз Одилов')

        cls.g1 = TeachingGroup.objects.create(
            class_subject=cls.cs_math, label='Гурӯҳи 1',
            teacher=cls.profile_a)
        cls.g2 = TeachingGroup.objects.create(
            class_subject=cls.cs_math, label='Гурӯҳи 2',
            teacher=cls.profile_a2)
        SubjectGroupMembership.objects.create(
            class_subject=cls.cs_math, group=cls.g1, student=cls.s1)
        SubjectGroupMembership.objects.create(
            class_subject=cls.cs_math, group=cls.g2, student=cls.s2)
        # s_extra deliberately has no membership (spec M).

    def _journal_url(self, school=None, class_name='7-А', subject=MATH):
        school = school or self.school_a
        return reverse('grade_entry', args=[school.id, class_name, subject])


class GroupAccessTests(GroupedFixture):
    """Entry-level access: who can open the grouped subject at all."""

    def test_group1_teacher_opens_journal(self):
        self.client.force_login(self.teacher_a)
        resp = self.client.get(self._journal_url())
        self.assertEqual(resp.status_code, 200)

    def test_group2_teacher_opens_journal_without_cs_assignment(self):
        # Regression for "two teachers on one subject, one cannot enter":
        # teacher_a2 has no ClassSubject.teacher/allocated_teacher — the
        # TeachingGroup assignment alone must grant journal access.
        self.client.force_login(self.teacher_a2)
        resp = self.client.get(self._journal_url())
        self.assertEqual(resp.status_code, 200)

    def test_group_only_teacher_reaches_class_detail(self):
        # The class picker only lists official curriculum subjects; the
        # grouped fixture subject (МАТЕМАТИКА) is not in TJC_SUBJECTS['7'],
        # so group the official 'ЗАБОНИ ТОҶИКӢ' row for this UI assertion.
        cs_tajik = make_class_subject(self.school_a, '7-А', TAJIK)
        g = TeachingGroup.objects.create(
            class_subject=cs_tajik, label='Гурӯҳи 1',
            teacher=self.profile_a2)
        SubjectGroupMembership.objects.create(
            class_subject=cs_tajik, group=g, student=self.s2)
        self.client.force_login(self.teacher_a2)
        resp = self.client.get(reverse(
            'class_detail', args=[self.school_a.id, '7-А']))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'ЗАБОНИ ТОҶИКӢ — Гурӯҳи 1')

    def test_group_only_teacher_sees_class_in_class_list(self):
        self.client.force_login(self.teacher_a2)
        resp = self.client.get(reverse('class_list', args=[self.school_a.id]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, '7-А')

    def test_group_only_teacher_does_not_see_unrelated_subject(self):
        # A grouped CS gives access to itself only — never to other subjects.
        make_class_subject(self.school_a, '7-А', TAJIK, teacher=None)
        self.client.force_login(self.teacher_a2)
        resp = self.client.get(
            self._journal_url(subject=TAJIK), follow=True)
        self.assertNotContains(resp, 'Холҳои ҳаррӯза')

    def test_cs_teacher_without_group_is_denied_on_grouped_subject(self):
        # Strict scope: a legacy CS-level assignment does NOT widen access
        # once active groups exist. teacher_x owns the CS row but no group.
        teacher_x = make_user('teacher_x', role='teacher', school=self.school_a)
        cs = make_class_subject(self.school_a, '7-А', TAJIK, teacher=teacher_x)
        TeachingGroup.objects.create(
            class_subject=cs, label='Гурӯҳи 1', teacher=self.profile_a2)
        TeachingGroup.objects.create(
            class_subject=cs, label='Гурӯҳи 2')
        self.client.force_login(teacher_x)
        resp = self.client.get(self._journal_url(subject=TAJIK))
        self.assertRedirects(
            resp, reverse('class_list', args=[self.school_a.id]))

    def test_admin_sees_both_groups(self):
        self.client.force_login(self.admin)
        resp = self.client.get(self._journal_url())
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, self.s1.full_name)
        self.assertContains(resp, self.s2.full_name)

    def test_zavuch_sees_both_groups(self):
        self.client.force_login(self.zavuch_a)
        resp = self.client.get(self._journal_url())
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, self.s1.full_name)
        self.assertContains(resp, self.s2.full_name)

    def test_cross_school_teacher_rejected(self):
        self.client.force_login(self.teacher_b)
        resp = self.client.get(self._journal_url())
        self.assertRedirects(
            resp, reverse('class_list', args=[self.school_b.id]))

    def test_teacher_group_badge_shown(self):
        self.client.force_login(self.teacher_a)
        resp = self.client.get(self._journal_url())
        self.assertContains(resp, 'Гурӯҳи 1')


class GroupStudentScopeTests(GroupedFixture):
    """A group teacher sees only their own group's students."""

    def test_group1_teacher_sees_only_group1_students(self):
        self.client.force_login(self.teacher_a)
        resp = self.client.get(self._journal_url())
        self.assertContains(resp, self.s1.full_name)
        self.assertNotContains(resp, self.s2.full_name)

    def test_group2_teacher_sees_only_group2_students(self):
        self.client.force_login(self.teacher_a2)
        resp = self.client.get(self._journal_url())
        self.assertContains(resp, self.s2.full_name)
        self.assertNotContains(resp, self.s1.full_name)

    def test_cs_level_teacher_does_not_widen_scope(self):
        # teacher_a is BOTH cs.teacher and g1 teacher — strict scope keeps
        # them inside Гурӯҳи 1 only.
        self.client.force_login(self.teacher_a)
        resp = self.client.get(self._journal_url())
        self.assertNotContains(resp, self.s2.full_name)

    def test_unassigned_student_hidden_from_group_teacher(self):
        self.client.force_login(self.teacher_a)
        resp = self.client.get(self._journal_url())
        self.assertNotContains(resp, self.s_extra.full_name)
        self.client.force_login(self.teacher_a2)
        resp = self.client.get(self._journal_url())
        self.assertNotContains(resp, self.s_extra.full_name)

    def test_unassigned_student_visible_to_admin(self):
        self.client.force_login(self.admin)
        resp = self.client.get(self._journal_url())
        self.assertContains(resp, self.s_extra.full_name)

    def test_monthly_journal_scoped(self):
        self.client.force_login(self.teacher_a)
        resp = self.client.get(reverse(
            'monthly_journal',
            args=[self.school_a.id, '7-А', MATH]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, self.s1.full_name)
        self.assertNotContains(resp, self.s2.full_name)

    def test_export_scoped_to_own_group(self):
        # The export writes one row per Grade record, so seed grades for
        # both groups' students before exporting as the group-1 teacher.
        make_grade(self.s1, subject=MATH, score=8, date=D_Q1)
        make_grade(self.s2, subject=MATH, score=9, date=D_Q1)
        self.client.force_login(self.teacher_a)
        resp = self.client.get(
            reverse('export_journal_excel',
                    args=[self.school_a.id, '7-А', MATH])
            + f'?month={D_Q1.strftime("%Y-%m")}')
        self.assertEqual(resp.status_code, 200)
        # XLSX is a zip — parse the sheet instead of byte-searching.
        ws = load_workbook(io.BytesIO(resp.content)).active
        names = [row[1] for row in ws.iter_rows(min_row=2, values_only=True)]
        self.assertIn(self.s1.full_name, names)
        self.assertNotIn(self.s2.full_name, names)


class GroupForgedWriteTests(GroupedFixture):
    """Server-side enforcement: forged IDs must be rejected."""

    def test_forged_student_id_ajax_rejected(self):
        self.client.force_login(self.teacher_a)
        resp = self.client.post(reverse('save_grade_ajax'), {
            'student_id': self.s2.id, 'subject': MATH,
            'type': 'daily', 'date': D_Q1.isoformat(), 'score': '8',
        })
        self.assertEqual(resp.status_code, 403)
        self.assertFalse(Grade.objects.filter(student=self.s2).exists())

    def test_forged_student_id_monthly_save_rejected(self):
        self.client.force_login(self.teacher_a)
        resp = self.client.post(reverse('monthly_save'), {
            'payload': json.dumps([{
                'key': 'k1', 'student_id': self.s2.id, 'subject': MATH,
                'date': D_Q1.isoformat(), 'score': '8',
            }]),
        })
        data = resp.json()
        self.assertFalse(data['results'][0]['ok'])
        self.assertFalse(Grade.objects.filter(student=self.s2).exists())

    def test_unassigned_student_write_rejected(self):
        self.client.force_login(self.teacher_a)
        resp = self.client.post(reverse('save_grade_ajax'), {
            'student_id': self.s_extra.id, 'subject': MATH,
            'type': 'daily', 'date': D_Q1.isoformat(), 'score': '8',
        })
        self.assertEqual(resp.status_code, 403)

    def test_forged_lesson_id_rejected(self):
        l_g2 = Lesson.objects.create(
            class_subject=self.cs_math, date=D_Q1,
            lesson_number=9, group=self.g2)
        self.client.force_login(self.teacher_a)
        # Own student but the other group's lesson — must be rejected.
        resp = self.client.post(reverse('save_grade_ajax'), {
            'student_id': self.s1.id, 'subject': MATH,
            'type': 'daily', 'date': D_Q1.isoformat(), 'score': '8',
            'lesson_id': str(l_g2.id),
        })
        data = resp.json()
        self.assertFalse(data['success'])
        self.assertFalse(Grade.objects.filter(student=self.s1).exists())

    def test_forged_assessment_id_rejected(self):
        l_g2 = Lesson.objects.create(
            class_subject=self.cs_math, date=D_Q1,
            lesson_number=9, group=self.g2)
        a2 = Assessment.objects.create(
            class_subject=self.cs_math, date=D_Q1, quarter=1,
            category='control', number=1, lesson=l_g2)
        self.client.force_login(self.teacher_a)
        resp = self.client.post(reverse('save_grade_ajax'), {
            'student_id': self.s1.id, 'subject': MATH,
            'type': 'daily', 'score': '8',
            'assessment_id': str(a2.id),
        })
        data = resp.json()
        self.assertFalse(data['success'])
        self.assertFalse(Grade.objects.filter(assessment=a2).exists())

    def test_batch_post_ignores_other_group_student(self):
        # The legacy batch POST iterates the scoped student list — a forged
        # score_<id> key for the other group is never read.
        self.client.force_login(self.teacher_a)
        resp = self.client.post(self._journal_url(), {
            'date': D_Q1.isoformat(),
            f'score_{self.s1.id}': '9',
            f'score_{self.s2.id}': '10',
        })
        self.assertEqual(resp.status_code, 302)
        self.assertTrue(Grade.objects.filter(student=self.s1).exists())
        self.assertFalse(Grade.objects.filter(student=self.s2).exists())

    def test_calc_quarter_scoped_to_own_group(self):
        Grade.objects.create(
            student=self.s1, subject=MATH, period='Холҳои ҷорӣ (Онлайн)',
            date=D_Q1, score=9)
        Grade.objects.create(
            student=self.s2, subject=MATH, period='Холҳои ҷорӣ (Онлайн)',
            date=D_Q1, score=5)
        self.client.force_login(self.teacher_a)
        self.client.post(reverse(
            'calc_quarter_from_daily',
            args=[self.school_a.id, '7-А', MATH]))
        from portal.models import QuarterGrade
        self.assertTrue(
            QuarterGrade.objects.filter(student=self.s1).exists())
        self.assertFalse(
            QuarterGrade.objects.filter(student=self.s2).exists())


class GroupLessonTests(GroupedFixture):
    """Lesson ownership: server-side group assignment, own-group only."""

    def _save(self, user, **post):
        self.client.force_login(user)
        return self.client.post(reverse('lesson_save'), post)

    def test_teacher_lesson_gets_own_group_server_side(self):
        resp = self._save(
            self.teacher_a,
            class_subject_id=str(self.cs_math.id),
            date=D_Q1.isoformat())
        self.assertTrue(resp.json()['success'])
        lesson = Lesson.objects.get(
            class_subject=self.cs_math, date=D_Q1)
        self.assertEqual(lesson.group_id, self.g1.id)

    def test_teacher_cannot_create_lesson_for_other_group(self):
        # A forged group_id is ignored — the server always assigns the
        # teacher's own group.
        resp = self._save(
            self.teacher_a,
            class_subject_id=str(self.cs_math.id),
            date=D_Q1.isoformat(), group_id=str(self.g2.id))
        self.assertTrue(resp.json()['success'])
        self.assertEqual(
            Lesson.objects.get(
                class_subject=self.cs_math, date=D_Q1).group_id,
            self.g1.id)

    def test_teacher_cannot_claim_other_group_lesson_slot(self):
        Lesson.objects.create(
            class_subject=self.cs_math, date=D_Q1,
            lesson_number=2, group=self.g2)
        resp = self._save(
            self.teacher_a,
            class_subject_id=str(self.cs_math.id),
            date=D_Q1.isoformat(), lesson_number='2', topic='x')
        self.assertEqual(resp.status_code, 403)

    def test_teacher_cannot_delete_other_group_lesson(self):
        l_g2 = Lesson.objects.create(
            class_subject=self.cs_math, date=D_Q1,
            lesson_number=2, group=self.g2)
        self.client.force_login(self.teacher_a)
        resp = self.client.post(reverse('lesson_delete'),
                                {'lesson_id': str(l_g2.id)})
        self.assertEqual(resp.status_code, 403)
        self.assertTrue(Lesson.objects.filter(pk=l_g2.pk).exists())

    def test_teacher_can_delete_own_empty_lesson(self):
        l_g1 = Lesson.objects.create(
            class_subject=self.cs_math, date=D_Q1,
            lesson_number=2, group=self.g1)
        self.client.force_login(self.teacher_a)
        resp = self.client.post(reverse('lesson_delete'),
                                {'lesson_id': str(l_g1.id)})
        self.assertTrue(resp.json()['success'])
        self.assertFalse(Lesson.objects.filter(pk=l_g1.pk).exists())

    def test_legacy_null_lesson_readable_by_both_groups(self):
        legacy = Lesson.objects.create(
            class_subject=self.cs_math, date=D_Q1, lesson_number=1)
        for teacher in (self.teacher_a, self.teacher_a2):
            self.client.force_login(teacher)
            resp = self.client.get(
                self._journal_url()
                + f'?date={D_Q1.isoformat()}&lesson={legacy.id}')
            self.assertEqual(resp.status_code, 200)

    def test_legacy_null_lesson_accepts_own_group_grade(self):
        legacy = Lesson.objects.create(
            class_subject=self.cs_math, date=D_Q1, lesson_number=1)
        self.client.force_login(self.teacher_a)
        resp = self.client.post(reverse('save_grade_ajax'), {
            'student_id': self.s1.id, 'subject': MATH,
            'type': 'daily', 'date': D_Q1.isoformat(), 'score': '7',
            'lesson_id': str(legacy.id),
        })
        self.assertTrue(resp.json()['success'])
        self.assertEqual(
            Grade.objects.get(student=self.s1).lesson_id, legacy.id)

    def test_other_group_lesson_hidden_from_teacher(self):
        Lesson.objects.create(
            class_subject=self.cs_math, date=D_Q1,
            lesson_number=1, group=self.g2)
        self.client.force_login(self.teacher_a)
        resp = self.client.get(
            self._journal_url() + f'?date={D_Q1.isoformat()}')
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, 'Дарси 1')

    def test_shared_numbering_preserved(self):
        # unique(class_subject, date, lesson_number) — group lessons share
        # the numbering space: G1 gets Дарси 1, G2 gets Дарси 2.
        self._save(self.teacher_a,
                   class_subject_id=str(self.cs_math.id),
                   date=D_Q1.isoformat())
        self._save(self.teacher_a2,
                   class_subject_id=str(self.cs_math.id),
                   date=D_Q1.isoformat())
        nums = list(
            Lesson.objects.filter(
                class_subject=self.cs_math, date=D_Q1)
            .order_by('lesson_number')
            .values_list('lesson_number', 'group_id'))
        self.assertEqual(nums, [(1, self.g1.id), (2, self.g2.id)])

    def test_admin_can_target_group_for_lesson(self):
        resp = self._save(
            self.admin,
            class_subject_id=str(self.cs_math.id),
            date=D_Q1.isoformat(), group_id=str(self.g2.id))
        self.assertTrue(resp.json()['success'])
        self.assertEqual(
            Lesson.objects.get(
                class_subject=self.cs_math, date=D_Q1).group_id,
            self.g2.id)


class GroupAssessmentTests(GroupedFixture):
    """Control assessments are group-owned via their lesson."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.l_g1 = Lesson.objects.create(
            class_subject=cls.cs_math, date=D_Q1,
            lesson_number=1, group=cls.g1)
        cls.l_g2 = Lesson.objects.create(
            class_subject=cls.cs_math, date=D_Q1,
            lesson_number=2, group=cls.g2)
        cls.a_g2 = Assessment.objects.create(
            class_subject=cls.cs_math, date=D_Q1, quarter=1,
            category='control', number=1, lesson=cls.l_g2)

    def _save(self, user, **post):
        self.client.force_login(user)
        return self.client.post(reverse('assessment_save'), post)

    def test_teacher_creates_assessment_on_own_lesson(self):
        resp = self._save(
            self.teacher_a,
            class_subject_id=str(self.cs_math.id),
            date=D_Q1.isoformat(), lesson_id=str(self.l_g1.id))
        self.assertTrue(resp.json()['success'])

    def test_teacher_lesson_less_assessment_rejected(self):
        resp = self._save(
            self.teacher_a,
            class_subject_id=str(self.cs_math.id),
            date=D_Q1.isoformat())
        self.assertFalse(resp.json()['success'])

    def test_teacher_cannot_attach_to_other_group_lesson(self):
        resp = self._save(
            self.teacher_a,
            class_subject_id=str(self.cs_math.id),
            date=D_Q1.isoformat(), lesson_id=str(self.l_g2.id))
        self.assertFalse(resp.json()['success'])

    def test_teacher_cannot_edit_other_group_assessment(self):
        resp = self._save(
            self.teacher_a,
            assessment_id=str(self.a_g2.id),
            date=D_Q1.isoformat(), lesson_id=str(self.l_g1.id),
            title='forged')
        self.assertFalse(resp.json()['success'])
        self.a_g2.refresh_from_db()
        self.assertNotEqual(self.a_g2.title, 'forged')

    def test_teacher_cannot_edit_lesson_less_assessment(self):
        a_admin = Assessment.objects.create(
            class_subject=self.cs_math, date=D_Q1, quarter=1,
            category='control', number=2)
        resp = self._save(
            self.teacher_a,
            assessment_id=str(a_admin.id),
            date=D_Q1.isoformat(), lesson_id=str(self.l_g1.id))
        self.assertFalse(resp.json()['success'])

    def test_teacher_cannot_delete_other_group_assessment(self):
        self.client.force_login(self.teacher_a)
        resp = self.client.post(reverse('assessment_delete'),
                                {'assessment_id': str(self.a_g2.id)})
        self.assertFalse(resp.json()['success'])
        self.assertTrue(Assessment.objects.filter(pk=self.a_g2.pk).exists())

    def test_admin_can_create_lesson_less_assessment(self):
        resp = self._save(
            self.admin,
            class_subject_id=str(self.cs_math.id),
            date=D_Q1.isoformat())
        self.assertTrue(resp.json()['success'])


class GroupLegacyTests(GoldenBase):
    """Ungrouped subjects and non-graded classes behave exactly as before."""

    def test_ungrouped_teacher_sees_whole_class(self):
        self.client.force_login(self.teacher_a)
        resp = self.client.get(reverse(
            'grade_entry', args=[self.school_a.id, '7-А', MATH]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, self.s1.full_name)
        self.assertContains(resp, self.s2.full_name)

    def test_ungrouped_lesson_save_unchanged(self):
        self.client.force_login(self.teacher_a)
        resp = self.client.post(reverse('lesson_save'), {
            'class_subject_id': str(self.cs_a.id),
            'date': D_Q1.isoformat()})
        self.assertTrue(resp.json()['success'])
        self.assertIsNone(Lesson.objects.get(
            class_subject=self.cs_a, date=D_Q1).group_id)

    def test_ungrouped_lesson_less_assessment_allowed(self):
        self.client.force_login(self.teacher_a)
        resp = self.client.post(reverse('assessment_save'), {
            'class_subject_id': str(self.cs_a.id),
            'date': D_Q1.isoformat()})
        self.assertTrue(resp.json()['success'])

    def test_ungrouped_write_unchanged(self):
        self.client.force_login(self.teacher_a)
        resp = self.client.post(reverse('save_grade_ajax'), {
            'student_id': self.s1.id, 'subject': MATH,
            'type': 'daily', 'date': D_Q1.isoformat(), 'score': '8'})
        self.assertTrue(resp.json()['success'])

    def test_non_graded_sticker_flow_unchanged(self):
        # 1-А is non-graded; a teacher assigned to any subject in it gets
        # the class-wide sticker journal — groups are ignored there.
        make_student(self.school_a, '1-А', 'Хонандаи Хурд')
        make_class_subject(
            self.school_a, '1-А', MATH, teacher=self.teacher_a)
        self.client.force_login(self.teacher_a)
        resp = self.client.get(reverse(
            'sticker_entry', args=[self.school_a.id, '1-А']))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Хонандаи Хурд')


class GroupRankingBaselineTests(GroupedFixture):
    """Ranking/activity formulas untouched: grouped grades count once."""

    def test_school_gpa_counts_each_student_once(self):
        from portal.views import calculate_school_rankings
        Grade.objects.create(
            student=self.s1, subject=MATH, period='Холҳои ҷорӣ (Онлайн)',
            date=D_Q1, score=8)
        Grade.objects.create(
            student=self.s2, subject=MATH, period='Холҳои ҷорӣ (Онлайн)',
            date=D_Q1, score=10)
        data = calculate_school_rankings()
        entry = next(d for d in data if d['school'].id == self.school_a.id)
        self.assertEqual(entry['gpa'], 9.0)
