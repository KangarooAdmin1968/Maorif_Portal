"""Ҳуҷҷатнигорӣ → Ҳисоботи маориф — district А.Ҷ.Ч. summary export.

One row per academic school: roster totals for levels 2-11, участие
(students with quarter grades), per-student quarter-average bands,
percent formulas, the Хамагӣ total row and the level-1 Эзоҳ note.
"""
import io

from django.urls import reverse
from openpyxl import load_workbook

from .helpers import (
    GoldenBase, MATH, make_quarter_grade, make_student, make_user,
)


class MaorifReportExportTests(GoldenBase):
    """GoldenBase: school A has s1+s2 in 7-А, school B has s3 in 7-Б."""

    FIRST = 6   # first data row in the generated sheet

    def _get(self, user, **params):
        self.client.force_login(user)
        return self.client.get(reverse('maorif_report_export'), params)

    def _ws(self, resp):
        return load_workbook(io.BytesIO(resp.content)).active

    def test_district_rows_bands_and_formulas(self):
        make_quarter_grade(self.s1, '7-А', MATH, 1, grade=9)
        make_quarter_grade(self.s3, '7-Б', MATH, 1, grade=3)
        resp = self._get(self.director)
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheetml', resp['Content-Type'])
        self.assertIn('attachment', resp['Content-Disposition'])
        ws = self._ws(resp)
        # Two school rows: МТМУ №1 then МТМУ №2.
        self.assertEqual(ws['B6'].value, 'МТМУ №1')
        self.assertEqual(ws['B7'].value, 'МТМУ №2')
        # School A: 2 rostered, 1 graded (s1 avg 9 -> 10-9 band).
        self.assertEqual(ws['C6'].value, 2)
        self.assertEqual(ws['D6'].value, 1)
        self.assertEqual(ws['E6'].value, 1)
        self.assertEqual(ws['G6'].value, 0)
        self.assertEqual(ws['K6'].value, 0)
        # School B: s3 avg 3 -> 3-1 band.
        self.assertEqual(ws['C7'].value, 1)
        self.assertEqual(ws['D7'].value, 1)
        self.assertEqual(ws['K7'].value, 1)
        # Percent / quality / mastery cells carry computed values — they
        # must display even in Excel's Protected View (no recalc there).
        self.assertEqual(ws['F6'].value, 100.0)   # 1/1 graded -> 10-9
        self.assertEqual(ws['M6'].value, 100.0)   # сифатнокӣ
        self.assertEqual(ws['N6'].value, 100.0)   # no 1-3 grades
        self.assertEqual(ws['N7'].value, 0.0)     # s3 fully in the 3-1 band
        # Хамагӣ total row holds summed counts.
        self.assertEqual(ws['B8'].value, 'Хамагӣ')
        self.assertEqual(ws['C8'].value, 3)
        self.assertEqual(ws['F8'].value, 50.0)    # 1 of 2 graded -> 10-9
        self.assertEqual(ws['M8'].value, 50.0)
        self.assertEqual(ws['N8'].value, 50.0)

    def test_ungraded_students_listed_in_note(self):
        # s1 gets a grade; s2 and s3 stay ungraded -> one row per school.
        make_quarter_grade(self.s1, '7-А', MATH, 1, grade=8)
        ws = self._ws(self._get(self.admin))
        note = ws['A10'].value or ''
        self.assertIn('синфи 1-ум', note)
        self.assertIn('Баҳо нагирифтаанд', ws['A11'].value or '')
        self.assertIn('МТМУ №1', ws['A12'].value or '')
        self.assertIn('7-А (1)', ws['A12'].value or '')
        self.assertIn('МТМУ №2', ws['A13'].value or '')
        self.assertIn('7-Б (1)', ws['A13'].value or '')

    def test_level1_students_only_in_note(self):
        make_student(self.school_a, '1-А', 'Сарвиноз Одилова')
        ws = self._ws(self._get(self.admin))
        self.assertEqual(ws['C6'].value, 2)   # 1-А excluded from the row
        self.assertIn('синфи 1-ум 1 нафар', ws['A10'].value)

    def test_zavuch_denied_district_data(self):
        resp = self._get(self.zavuch_a)
        self.assertEqual(resp.status_code, 302)

    def test_teacher_denied(self):
        resp = self._get(self.teacher_a)
        self.assertEqual(resp.status_code, 302)

    def test_staff_and_director_allowed(self):
        staff = make_user('maorif_staff', staff=True)
        for user in (self.director, staff):
            resp = self._get(user)
            self.assertEqual(resp.status_code, 200, user.username)

    def test_director_reaches_documents_page(self):
        # The director now sees Ҳуҷҷатнигорӣ with the district button.
        self.client.force_login(self.director)
        resp = self.client.get(reverse('school_documents'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, reverse('maorif_reports'))

    def test_district_users_pick_school_on_documents_page(self):
        # Маориф staff and the district director get a school selector;
        # once a school is chosen every document link carries school_id.
        staff = make_user('maorif_staff_docs', staff=True)
        for user in (self.director, staff):
            self.client.force_login(user)
            resp = self.client.get(reverse('school_documents'))
            self.assertEqual(resp.status_code, 200, user.username)
            self.assertContains(resp, 'name="school_id"')
            resp = self.client.get(
                reverse('school_documents'),
                {'school_id': self.school_a.id})
            self.assertEqual(resp.status_code, 200)
            self.assertEqual(resp.context['school'], self.school_a)
            self.assertContains(
                resp, f'school_id={self.school_a.id}')
            self.assertContains(resp, 'Ҳисоботи завуч')

    def test_district_users_download_school_reports(self):
        # The per-school exports that used to kick district users back to
        # the dashboard now accept ?school_id= — read-only access.
        staff = make_user('maorif_staff_zr', staff=True)
        for user in (self.director, staff):
            self.client.force_login(user)
            for url_name in ('zavuch_report_export',
                             'zavuch_grades_export'):
                resp = self.client.get(reverse(url_name), {
                    'quarter': 1, 'school_id': self.school_a.id})
                self.assertEqual(resp.status_code, 200,
                                 f'{user.username} {url_name}')
                self.assertIn('attachment', resp['Content-Disposition'])

    def test_reports_page_lists_files(self):
        self.client.force_login(self.director)
        resp = self.client.get(reverse('maorif_reports'), {'quarter': 2})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, reverse('maorif_report_export'))
        self.assertContains(resp, reverse('maorif_report_total_export'))

    def test_total_workbook_single_khamagi_row(self):
        make_quarter_grade(self.s1, '7-А', MATH, 1, grade=9)
        make_quarter_grade(self.s3, '7-Б', MATH, 1, grade=3)
        self.client.force_login(self.admin)
        resp = self.client.get(reverse('maorif_report_total_export'))
        self.assertEqual(resp.status_code, 200)
        ws = load_workbook(io.BytesIO(resp.content)).active
        # Single bold Хамагӣ row as the only data row.
        self.assertEqual(ws['A6'].value, 1)
        self.assertEqual(ws['B6'].value, 'Хамагӣ')
        self.assertEqual(ws['C6'].value, 3)    # s1+s2+s3
        self.assertEqual(ws['D6'].value, 2)    # s1, s3 graded
        self.assertEqual(ws['E6'].value, 1)    # s1 -> 10-9
        self.assertEqual(ws['K6'].value, 1)    # s3 -> 3-1
        self.assertIsNone(ws['B7'].value)      # no second data row

    def test_reports_page_denied_for_zavuch(self):
        self.client.force_login(self.zavuch_a)
        resp = self.client.get(reverse('maorif_reports'))
        self.assertEqual(resp.status_code, 302)
        resp = self.client.get(reverse('maorif_report_total_export'))
        self.assertEqual(resp.status_code, 302)

    def test_grades_workbook_per_mark_columns(self):
        # s1 avg 10 -> mark column "10"; s3 avg 3 -> mark column "3".
        make_quarter_grade(self.s1, '7-А', MATH, 1, grade=10)
        make_quarter_grade(self.s3, '7-Б', MATH, 1, grade=3)
        self.client.force_login(self.admin)
        resp = self.client.get(reverse('maorif_report_grades_export'))
        self.assertEqual(resp.status_code, 200)
        ws = self._ws(resp)
        # Sub-header row carries each mark + % pair, 10 down to 1.
        self.assertEqual(ws['E5'].value, '10')
        self.assertEqual(ws['F5'].value, '%')
        self.assertEqual(ws['W5'].value, '1')
        # School A row: s1 in the "10" column (E=count, F=%).
        self.assertEqual(ws['B6'].value, 'МТМУ №1')
        self.assertEqual(ws['E6'].value, 1)
        self.assertEqual(ws['F6'].value, 100.0)
        self.assertEqual(ws['Y6'].value, 100.0)  # сифат 7-10
        self.assertEqual(ws['Z6'].value, 100.0)  # азхудкунӣ 4-10
        # School B row: s3 in the "3" column (S=count, T=%).
        self.assertEqual(ws['S7'].value, 1)
        self.assertEqual(ws['Y7'].value, 0.0)
        self.assertEqual(ws['Z7'].value, 0.0)
        # Хамагӣ row follows the two school rows.
        self.assertEqual(ws['B8'].value, 'Хамагӣ')
        self.assertEqual(ws['E8'].value, 1)
        self.assertEqual(ws['S8'].value, 1)

    def test_grades_district_workbook_single_row(self):
        make_quarter_grade(self.s1, '7-А', MATH, 1, grade=10)
        self.client.force_login(self.director)
        resp = self.client.get(
            reverse('maorif_report_grades_district_export'))
        self.assertEqual(resp.status_code, 200)
        ws = self._ws(resp)
        # Official one-row form: «ноҳияи Зафаробод» + Хамагӣ.
        self.assertEqual(ws['A6'].value, 1)
        self.assertEqual(ws['B6'].value, 'ноҳияи Зафаробод')
        self.assertEqual(ws['D6'].value, 1)
        self.assertEqual(ws['E6'].value, 1)
        self.assertEqual(ws['B7'].value, 'Хамагӣ')
        self.assertEqual(ws['E7'].value, 1)
        self.assertIsNone(ws['B8'].value)

    def test_reports_page_lists_grade_files(self):
        self.client.force_login(self.director)
        resp = self.client.get(reverse('maorif_reports'))
        self.assertContains(resp, reverse('maorif_report_grades_export'))
        self.assertContains(
            resp, reverse('maorif_report_grades_district_export'))
