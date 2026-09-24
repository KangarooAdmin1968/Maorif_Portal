"""Tests for the monitoring dashboard + Excel export permission scoping.

Covers the Phase 4 wiring of portal.monitoring helpers into
monitoring_dashboard and export_monitoring_excel: district viewers keep
full detail, principals see only their own school, zavuchs see aggregate
rows for other schools with no individual usernames/last_login data.
"""
import io
import datetime

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from openpyxl import load_workbook

from portal.tests.helpers import make_school, make_user


class MonitoringDashboardBase(TestCase):

    @classmethod
    def setUpTestData(cls):
        cls.school_a = make_school('Мактаб №1')
        cls.school_b = make_school('Мактаб №2')
        cls.admin = make_user('admin_test', superuser=True)
        cls.staff_u = make_user('staff_1', staff=True)
        cls.director = make_user('director_1', role='director')
        cls.principal_a = make_user('principal_1', role='principal',
                                  school=cls.school_a)
        cls.zavuch_a = make_user('zavuch_1', role='teacher',
                                 school=cls.school_a)
        cls.zavuch_b = make_user('zavuch_2', role='teacher',
                                 school=cls.school_b)
        cls.teacher_a = make_user('teacher_1', role='teacher',
                                  school=cls.school_a)
        cls.bare = make_user('plain_user')
        # Give school B's zavuch a deterministic last_login to mask-test.
        cls.zavuch_b.last_login = timezone.make_aware(
            datetime.datetime(2026, 9, 20, 10, 30))
        cls.zavuch_b.save(update_fields=['last_login'])

    def _get_dashboard(self, user):
        self.client.force_login(user)
        return self.client.get(reverse('monitoring_dashboard'))

    def _get_export(self, user):
        self.client.force_login(user)
        return self.client.get(reverse('export_monitoring_excel'))

    def _row_for(self, resp, school):
        for row in resp.context['stats']:
            if row['school'].id == school.id:
                return row
        return None

    def _workbook_rows(self, resp):
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        rows = []
        for r in range(2, ws.max_row + 1):
            rows.append([ws.cell(row=r, column=c).value for c in range(1, 8)])
        return rows

    def _other_school_login_str(self):
        return timezone.localtime(self.zavuch_b.last_login).strftime(
            '%d.%m.%Y %H:%M')


class DashboardDistrictScopeTests(MonitoringDashboardBase):

    def _assert_full_detail(self, user):
        resp = self._get_dashboard(user)
        self.assertEqual(resp.status_code, 200)
        schools = {r['school'].id for r in resp.context['stats']}
        self.assertEqual(schools, {self.school_a.id, self.school_b.id})
        row_b = self._row_for(resp, self.school_b)
        self.assertFalse(row_b.get('masked'))
        self.assertEqual(row_b['zavuch_username'], 'zavuch_2')
        self.assertEqual(row_b['zavuch_user'], self.zavuch_b)
        self.assertTrue(row_b['logged_in'])
        # Existing calculations remain intact for full-detail viewers.
        for key in ('class_count', 'teacher_count', 'student_count'):
            self.assertIn(key, row_b)

    def test_superuser_sees_all_schools_full_detail(self):
        self._assert_full_detail(self.admin)

    def test_staff_sees_all_schools_full_detail(self):
        self._assert_full_detail(self.staff_u)

    def test_district_director_sees_all_schools_full_detail(self):
        self._assert_full_detail(self.director)


class DashboardPrincipalScopeTests(MonitoringDashboardBase):

    def test_principal_sees_only_own_school(self):
        resp = self._get_dashboard(self.principal_a)
        self.assertEqual(resp.status_code, 200)
        stats = resp.context['stats']
        self.assertEqual(len(stats), 1)
        self.assertEqual(stats[0]['school'].id, self.school_a.id)

    def test_principal_cannot_receive_other_school_data(self):
        resp = self._get_dashboard(self.principal_a)
        self.assertIsNone(self._row_for(resp, self.school_b))
        self.assertNotContains(resp, self.school_b.name)
        self.assertNotContains(resp, 'zavuch_2')


class DashboardZavuchScopeTests(MonitoringDashboardBase):

    def test_zavuch_sees_own_school_full_detail(self):
        resp = self._get_dashboard(self.zavuch_a)
        self.assertEqual(resp.status_code, 200)
        row_a = self._row_for(resp, self.school_a)
        self.assertFalse(row_a.get('masked'))
        self.assertEqual(row_a['zavuch_username'], 'zavuch_1')
        self.assertEqual(row_a['zavuch_user'], self.zavuch_a)

    def test_zavuch_other_schools_aggregate_only(self):
        resp = self._get_dashboard(self.zavuch_a)
        row_b = self._row_for(resp, self.school_b)
        self.assertIsNotNone(row_b)
        self.assertTrue(row_b['masked'])
        self.assertIsNone(row_b['zavuch_user'])
        # Aggregate columns stay visible.
        self.assertIn('class_count', row_b)
        self.assertIn('student_count', row_b)

    def test_zavuch_response_has_no_other_school_username(self):
        resp = self._get_dashboard(self.zavuch_a)
        self.assertNotContains(resp, 'zavuch_2')

    def test_zavuch_response_has_no_other_school_last_login(self):
        resp = self._get_dashboard(self.zavuch_a)
        self.assertNotContains(resp, self._other_school_login_str())


class DashboardDeniedTests(MonitoringDashboardBase):

    def test_regular_teacher_denied(self):
        resp = self._get_dashboard(self.teacher_a)
        self.assertRedirects(resp, reverse('dashboard'),
                             fetch_redirect_response=False)

    def test_profileless_user_denied(self):
        resp = self._get_dashboard(self.bare)
        self.assertRedirects(resp, reverse('dashboard'),
                             fetch_redirect_response=False)

    def test_anonymous_denied(self):
        resp = self.client.get(reverse('monitoring_dashboard'))
        self.assertEqual(resp.status_code, 302)
        self.assertIn('/login/', resp.url)


class ExportScopeTests(MonitoringDashboardBase):

    def _assert_all_schools_export(self, user):
        resp = self._get_export(user)
        self.assertEqual(resp.status_code, 200)
        rows = self._workbook_rows(resp)
        names = {r[1] for r in rows}
        self.assertEqual(names, {self.school_a.name, self.school_b.name})
        usernames = {r[2] for r in rows}
        self.assertIn('zavuch_2', usernames)

    def test_superuser_exports_all_schools(self):
        self._assert_all_schools_export(self.admin)

    def test_staff_exports_all_schools(self):
        self._assert_all_schools_export(self.staff_u)

    def test_director_exports_all_schools(self):
        self._assert_all_schools_export(self.director)

    def test_principal_exports_only_own_school(self):
        resp = self._get_export(self.principal_a)
        self.assertEqual(resp.status_code, 200)
        rows = self._workbook_rows(resp)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][1], self.school_a.name)
        self.assertEqual(rows[0][2], 'zavuch_1')

    def test_zavuch_export_own_school_detail(self):
        resp = self._get_export(self.zavuch_a)
        rows = self._workbook_rows(resp)
        own = next(r for r in rows if r[1] == self.school_a.name)
        self.assertEqual(own[2], 'zavuch_1')

    def test_zavuch_export_other_schools_aggregate_only(self):
        resp = self._get_export(self.zavuch_a)
        rows = self._workbook_rows(resp)
        other = next(r for r in rows if r[1] == self.school_b.name)
        self.assertEqual(other[2], '—')   # username column masked
        self.assertEqual(other[3], '—')   # login column masked
        self.assertIsInstance(other[4], int)  # aggregate counts remain
        self.assertIsInstance(other[6], int)

    def test_zavuch_export_has_no_other_school_username(self):
        resp = self._get_export(self.zavuch_a)
        rows = self._workbook_rows(resp)
        flat = [str(cell) for row in rows for cell in row]
        self.assertFalse(any('zavuch_2' in cell for cell in flat))

    def test_zavuch_export_has_no_other_school_login_status(self):
        resp = self._get_export(self.zavuch_a)
        rows = self._workbook_rows(resp)
        other = next(r for r in rows if r[1] == self.school_b.name)
        self.assertNotIn(other[3], ('Фаъол', 'Ғайрифаъол'))

    def test_teacher_export_denied(self):
        resp = self._get_export(self.teacher_a)
        self.assertRedirects(resp, reverse('dashboard'),
                             fetch_redirect_response=False)

    def test_profileless_export_denied(self):
        resp = self._get_export(self.bare)
        self.assertRedirects(resp, reverse('dashboard'),
                             fetch_redirect_response=False)

    def test_anonymous_export_denied(self):
        resp = self.client.get(reverse('export_monitoring_excel'))
        self.assertEqual(resp.status_code, 302)
        self.assertIn('/login/', resp.url)
