"""Stage 4 tests: monthly journal view.

Covers: month grid shape, per-lesson columns (no same-day collapse),
legacy lesson-less slot visibility, completion statuses, control markers,
month navigation, permissions, and the mobile day-list markup.
"""
import datetime
import json

from django.urls import reverse

from portal.models import Assessment, Grade
from portal.tests.helpers import D_Q1, D_Q1_B, MATH, GoldenBase
from portal.tests.test_phase4_lessons import (
    make_control, make_lesson, post_grade,
)

OCT = '2025-10'
NOV = '2025-11'
D_OCT_A = datetime.date(2025, 10, 15)   # D_Q1
D_OCT_B = datetime.date(2025, 10, 20)
D_OCT_C = datetime.date(2025, 10, 22)


def col_index(ctx, date, lesson_id):
    for i, col in enumerate(ctx['columns']):
        if col['date'] == date and col['lesson_id'] == lesson_id:
            return i
    return None


def cells_for(ctx, student, date, lesson_id):
    idx = col_index(ctx, date, lesson_id)
    assert idx is not None, f'no column for {date} lesson={lesson_id}'
    row = next(r for r in ctx['table_rows'] if r['student'].id == student.id)
    return row['cells'][idx]['items']


class MonthlyJournalTests(GoldenBase):
    def setUp(self):
        super().setUp()
        self.client.force_login(self.teacher_a)
        self.url = reverse('monthly_journal', kwargs={
            'school_id': self.school_a.id, 'class_name': '7-А', 'subject': MATH,
        })

    def _get(self, **params):
        params.setdefault('month', OCT)
        return self.client.get(self.url, params)

    def _col_index(self, ctx, date, lesson_id):
        return col_index(ctx, date, lesson_id)

    def _cells_for(self, ctx, student, date, lesson_id):
        return cells_for(ctx, student, date, lesson_id)

    # -- grid shape --------------------------------------------------------

    def test_get_renders_200(self):
        resp = self._get()
        self.assertEqual(resp.status_code, 200)
        self.assertTemplateUsed(resp, 'portal/monthly_journal.html')

    def test_all_days_present(self):
        resp = self._get()
        self.assertEqual(len(resp.context['header_groups']), 31)
        self.assertEqual(len(resp.context['day_rows']), 31)

    def test_one_column_per_lesson_same_day(self):
        l1 = make_lesson(self.cs_a, D_OCT_A, 1)
        l2 = make_lesson(self.cs_a, D_OCT_A, 2)
        resp = self._get()
        idx1 = self._col_index(resp.context, D_OCT_A, l1.id)
        idx2 = self._col_index(resp.context, D_OCT_A, l2.id)
        self.assertIsNotNone(idx1)
        self.assertIsNotNone(idx2)
        # No extra legacy column when it holds no data.
        self.assertIsNone(self._col_index(resp.context, D_OCT_A, None))

    def test_same_day_lesson_grades_not_collapsed(self):
        l1 = make_lesson(self.cs_a, D_OCT_A, 1)
        l2 = make_lesson(self.cs_a, D_OCT_A, 2)
        post_grade(self.client, self.s1, score='8', lesson=l1)
        post_grade(self.client, self.s1, score='10', lesson=l2)
        resp = self._get()
        c1 = self._cells_for(resp.context, self.s1, D_OCT_A, l1.id)
        c2 = self._cells_for(resp.context, self.s1, D_OCT_A, l2.id)
        self.assertEqual([c['text'] for c in c1], ['8'])
        self.assertEqual([c['text'] for c in c2], ['10'])

    def test_legacy_slot_column_only_when_data(self):
        make_lesson(self.cs_a, D_OCT_A, 1)
        post_grade(self.client, self.s1, score='9')  # lesson-less legacy row
        resp = self._get()
        idx = self._col_index(resp.context, D_OCT_A, None)
        self.assertIsNotNone(idx)
        cells = self._cells_for(resp.context, self.s1, D_OCT_A, None)
        self.assertEqual([c['text'] for c in cells], ['9'])

    def test_day_without_lesson_single_column(self):
        post_grade(self.client, self.s1, score='7', date=D_OCT_B)
        resp = self._get()
        idx = self._col_index(resp.context, D_OCT_B, None)
        self.assertIsNotNone(idx)

    # -- status derivation --------------------------------------------------

    def test_status_empty_partial_complete(self):
        post_grade(self.client, self.s1, score='7', date=D_OCT_B)
        post_grade(self.client, self.s1, score='8', date=D_OCT_C)
        post_grade(self.client, self.s2, score='9', date=D_OCT_C)
        resp = self._get()
        status = {d['date']: d['status'] for d in resp.context['day_rows']}
        self.assertEqual(status[D_OCT_A], 'empty')
        self.assertEqual(status[D_OCT_B], 'partial')
        self.assertEqual(status[D_OCT_C], 'complete')

    def test_control_flag_on_date_and_column(self):
        l1 = make_lesson(self.cs_a, D_OCT_A, 1)
        l2 = make_lesson(self.cs_a, D_OCT_A, 2)
        make_control(self.cs_a, date=D_OCT_A, lesson=l1)
        resp = self._get()
        day = next(d for d in resp.context['day_rows'] if d['date'] == D_OCT_A)
        self.assertTrue(day['has_control'])
        self.assertTrue(day['multi_lesson'])
        idx1 = self._col_index(resp.context, D_OCT_A, l1.id)
        idx2 = self._col_index(resp.context, D_OCT_A, l2.id)
        self.assertTrue(resp.context['columns'][idx1]['control'])
        self.assertFalse(resp.context['columns'][idx2]['control'])

    def test_control_grade_cell_marked(self):
        a = make_control(self.cs_a, date=D_OCT_A)
        post_grade(self.client, self.s1, score='6', assessment=a)
        resp = self._get()
        cells = self._cells_for(resp.context, self.s1, D_OCT_A, None)
        self.assertEqual([c['text'] for c in cells], ['6'])
        self.assertTrue(cells[0]['control'])

    def test_absent_mark_shown(self):
        post_grade(self.client, self.s1, extra={'attendance': '-'})
        resp = self._get()
        cells = self._cells_for(resp.context, self.s1, D_OCT_A, None)
        self.assertEqual([c['text'] for c in cells], ['н'])
        self.assertTrue(cells[0]['absent'])

    def test_paired_work_shows_slash_not_mean(self):
        a = make_control(self.cs_a, date=D_OCT_A)
        a.work_type = 'retelling'
        a.save(update_fields=['work_type'])
        post_grade(self.client, self.s1, score='8/9', assessment=a)
        resp = self._get()
        cells = self._cells_for(resp.context, self.s1, D_OCT_A, None)
        self.assertEqual([c['text'] for c in cells], ['8/9'])

    # -- month navigation ---------------------------------------------------

    def test_month_param_selects_month(self):
        post_grade(self.client, self.s1, score='7', date=D_Q1_B)  # November
        resp = self._get(month=OCT)
        # Nov 20 is not in the October grid at all.
        self.assertEqual(len(resp.context['day_rows']), 31)
        self.assertNotIn(
            D_Q1_B, [d['date'] for d in resp.context['day_rows']]
        )
        resp = self._get(month=NOV)
        status = {d['date']: d['status'] for d in resp.context['day_rows']}
        self.assertEqual(status[D_Q1_B], 'partial')
        self.assertEqual(len(resp.context['day_rows']), 30)

    def test_prev_next_month_links(self):
        resp = self._get(month=OCT)
        self.assertEqual(resp.context['prev_month'], datetime.date(2025, 9, 1))
        self.assertEqual(resp.context['next_month'], datetime.date(2025, 11, 1))

    def test_invalid_month_falls_back_to_current(self):
        resp = self.client.get(self.url, {'month': 'not-a-month'})
        self.assertEqual(resp.status_code, 200)
        today = datetime.date.today()
        self.assertEqual(
            resp.context['month_start'], today.replace(day=1)
        )

    # -- permissions ---------------------------------------------------------

    def test_cross_school_teacher_redirected(self):
        self.client.force_login(self.teacher_b)
        resp = self._get()
        self.assertEqual(resp.status_code, 302)

    def test_wrong_subject_teacher_forbidden_or_redirect(self):
        self.client.force_login(self.teacher_a)
        url = reverse('monthly_journal', kwargs={
            'school_id': self.school_a.id, 'class_name': '7-А',
            'subject': 'ЗАБОНИ ТОҶИКӢ',
        })
        resp = self.client.get(url, {'month': OCT})
        self.assertIn(resp.status_code, (302, 403))

    def test_anonymous_redirected(self):
        self.client.logout()
        resp = self._get()
        self.assertEqual(resp.status_code, 302)

    # -- mobile-safe rendering ------------------------------------------------

    def test_mobile_day_list_rendered(self):
        l1 = make_lesson(self.cs_a, D_OCT_A, 1)
        resp = self._get()
        html = resp.content.decode()
        self.assertIn('day-list', html)
        self.assertIn('month-wrap', html)
        # Lesson chip links into the daily journal for that lesson.
        self.assertIn(f'lesson={l1.id}', html)


class JournalExcelExportTests(GoldenBase):
    """Stage 6: flat lesson-aware export over the same journal data."""

    def setUp(self):
        super().setUp()
        self.client.force_login(self.teacher_a)
        self.url = reverse('export_journal_excel', kwargs={
            'school_id': self.school_a.id, 'class_name': '7-А', 'subject': MATH,
        })

    def _rows(self, resp):
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        return [tuple(c.value for c in r) for r in ws.iter_rows()]

    def test_export_returns_xlsx(self):
        resp = self.client.get(self.url, {'month': OCT})
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheetml', resp['Content-Type'])
        self.assertIn('attachment', resp['Content-Disposition'])

    def test_lesson_number_preserved_per_row(self):
        l1 = make_lesson(self.cs_a, D_OCT_A, 1)
        l2 = make_lesson(self.cs_a, D_OCT_A, 2)
        post_grade(self.client, self.s1, score='8', lesson=l1)
        post_grade(self.client, self.s1, score='10', lesson=l2)
        resp = self.client.get(self.url, {'month': OCT})
        rows = self._rows(resp)
        self.assertEqual(rows[0][:4], ('№', 'Хонанда', 'Сана', 'Дарс'))
        data = [r for r in rows[1:] if r[1] == self.s1.full_name]
        self.assertEqual(len(data), 2)
        self.assertEqual(
            sorted((r[3], r[6]) for r in data), [(1, '8'), (2, '10')]
        )

    def test_legacy_row_exports_empty_lesson(self):
        post_grade(self.client, self.s1, score='9')
        resp = self.client.get(self.url, {'month': OCT})
        rows = self._rows(resp)
        data = [r for r in rows[1:] if r[1] == self.s1.full_name]
        self.assertEqual(len(data), 1)
        self.assertIn(data[0][3], (None, ''))
        self.assertEqual(data[0][6], '9')
        self.assertEqual(data[0][4], 'Ҷорӣ')

    def test_control_row_marked(self):
        a = make_control(self.cs_a, date=D_OCT_A, title='Санҷиши 1')
        post_grade(self.client, self.s1, score='7', assessment=a)
        resp = self.client.get(self.url, {'month': OCT})
        rows = self._rows(resp)
        data = [r for r in rows[1:] if r[1] == self.s1.full_name]
        self.assertEqual(data[0][4], 'Назорат')
        self.assertIn('Санҷиши 1', data[0][5])

    def test_out_of_month_grades_excluded(self):
        post_grade(self.client, self.s1, score='7', date=D_Q1_B)  # Nov
        post_grade(self.client, self.s1, score='8', date=D_OCT_A)
        resp = self.client.get(self.url, {'month': OCT})
        rows = self._rows(resp)
        self.assertEqual(len(rows), 2)  # header + one grade row

    def test_cross_school_teacher_denied(self):
        self.client.force_login(self.teacher_b)
        resp = self.client.get(self.url, {'month': OCT})
        self.assertIn(resp.status_code, (302, 403))


# ---------------------------------------------------------------------------
# monthly_save — batch write endpoint for the editable monthly journal
# ---------------------------------------------------------------------------

MONTHLY_SAVE_URL = reverse('monthly_save')


class MonthlySaveTests(GoldenBase):
    def setUp(self):
        super().setUp()
        self.client.force_login(self.teacher_a)

    def _post(self, items):
        return self.client.post(
            MONTHLY_SAVE_URL, {'payload': json.dumps(items)}
        )

    def _item(self, student=None, date=D_OCT_A, lesson=None, **fields):
        item = {
            'key': f'{student.id if student else self.s1.id}|{date.isoformat()}|{lesson or ""}',
            'student_id': student.id if student else self.s1.id,
            'subject': MATH,
            'date': date.isoformat(),
        }
        if lesson:
            item['lesson_id'] = lesson
        item.update(fields)
        return item

    def test_create_grade(self):
        resp = self._post([self._item(score='8')])
        self.assertTrue(resp.json()['success'])
        self.assertEqual(resp.json()['saved'], 1)
        g = Grade.objects.get(student=self.s1, date=D_OCT_A)
        self.assertEqual(g.score, 8)
        self.assertIsNone(g.lesson_id)
        self.assertIsNone(g.assessment_id)

    def test_edit_existing_grade(self):
        post_grade(self.client, self.s1, score='7', date=D_OCT_A)
        resp = self._post([self._item(score='9')])
        self.assertTrue(resp.json()['success'])
        g = Grade.objects.get(student=self.s1, date=D_OCT_A)
        self.assertEqual(g.score, 9)

    def test_batch_multiple_cells_one_request(self):
        resp = self._post([
            self._item(self.s1, D_OCT_A, score='8'),
            self._item(self.s2, D_OCT_A, score='9'),
            self._item(self.s1, D_OCT_B, score='7'),
        ])
        self.assertEqual(resp.json()['saved'], 3)
        self.assertEqual(Grade.objects.filter(subject=MATH).count(), 3)

    def test_clear_score_deletes_empty_row(self):
        post_grade(self.client, self.s1, score='7', date=D_OCT_A)
        resp = self._post([self._item(score='')])
        self.assertTrue(resp.json()['success'])
        self.assertFalse(Grade.objects.filter(student=self.s1).exists())

    def test_lesson_scoped_write(self):
        l1 = make_lesson(self.cs_a, D_OCT_A, 1)
        l2 = make_lesson(self.cs_a, D_OCT_A, 2)
        resp = self._post([
            self._item(lesson=l1.id, score='8'),
            self._item(lesson=l2.id, score='10'),
        ])
        self.assertEqual(resp.json()['saved'], 2)
        self.assertEqual(
            Grade.objects.get(student=self.s1, lesson=l1).score, 8
        )
        self.assertEqual(
            Grade.objects.get(student=self.s1, lesson=l2).score, 10
        )

    def test_attendance_and_behavior_write(self):
        resp = self._post([self._item(attendance='-', behavior_score='4')])
        self.assertEqual(resp.json()['saved'], 1)
        g = Grade.objects.get(student=self.s1, date=D_OCT_A)
        self.assertEqual(g.attendance, '-')
        self.assertEqual(g.behavior_score, 4)

    def test_invalid_score_rejected_item_only(self):
        resp = self._post([
            self._item(self.s1, D_OCT_A, score='99'),
            self._item(self.s2, D_OCT_A, score='7'),
        ])
        data = resp.json()
        self.assertEqual(data['saved'], 1)
        self.assertEqual(
            Grade.objects.get(student=self.s2, date=D_OCT_A).score, 7
        )
        # The invalid row is reported per-item.
        bad = next(r for r in data['results'] if not r['ok'])
        self.assertIn('Хол', bad['message'])

    def test_quarter_lock_blocks_item(self):
        from portal.tests.helpers import lock_quarter
        lock_quarter(self.school_a, '7-А', MATH, quarter=1)
        resp = self._post([self._item(score='8')])
        data = resp.json()
        self.assertEqual(data['saved'], 0)
        self.assertIn('баста', data['results'][0]['message'])
        self.assertFalse(Grade.objects.filter(student=self.s1).exists())

    def test_control_slot_blocks_regular_write(self):
        l1 = make_lesson(self.cs_a, D_OCT_A, 1)
        make_control(self.cs_a, date=D_OCT_A, lesson=l1)
        resp = self._post([self._item(lesson=l1.id, score='8')])
        self.assertEqual(resp.json()['saved'], 0)
        self.assertFalse(
            Grade.objects.filter(student=self.s1, lesson=l1).exists()
        )

    def test_control_lesson_does_not_block_other_lesson(self):
        l1 = make_lesson(self.cs_a, D_OCT_A, 1)
        l2 = make_lesson(self.cs_a, D_OCT_A, 2)
        make_control(self.cs_a, date=D_OCT_A, lesson=l1)
        resp = self._post([self._item(lesson=l2.id, score='8')])
        self.assertEqual(resp.json()['saved'], 1)

    def test_foreign_lesson_rejected(self):
        foreign = make_lesson(self.cs_b, D_OCT_A, 1)
        resp = self._post([self._item(lesson=foreign.id, score='8')])
        self.assertEqual(resp.json()['saved'], 0)
        self.assertFalse(Grade.objects.filter(student=self.s1).exists())

    def test_foreign_student_rejected(self):
        # s3 belongs to School B — teacher_a must not write it.
        resp = self._post([self._item(student=self.s3, score='8')])
        data = resp.json()
        self.assertEqual(data['saved'], 0)
        self.assertFalse(
            Grade.objects.filter(student=self.s3).exists()
        )

    def test_malformed_payload(self):
        resp = self.client.post(MONTHLY_SAVE_URL, {'payload': '{oops'})
        self.assertFalse(resp.json()['success'])

    def test_unauthenticated_rejected(self):
        self.client.logout()
        resp = self._post([self._item(score='8')])
        self.assertEqual(resp.status_code, 302)

    # -- daily ↔ monthly synchronization (one source of truth) ---------------

    def test_monthly_write_visible_in_daily(self):
        self._post([self._item(score='8')])
        resp = self.client.get(reverse('grade_entry', kwargs={
            'school_id': self.school_a.id, 'class_name': '7-А', 'subject': MATH,
        }), {'date': D_OCT_A.isoformat()})
        self.assertEqual(resp.context['daily_grades'].get(self.s1.id), 8)

    def test_daily_write_visible_in_monthly(self):
        post_grade(self.client, self.s1, score='6', date=D_OCT_A)
        resp = self.client.get(reverse('monthly_journal', kwargs={
            'school_id': self.school_a.id, 'class_name': '7-А', 'subject': MATH,
        }), {'month': OCT})
        cells = cells_for(resp.context, self.s1, D_OCT_A, None)
        self.assertEqual([c['text'] for c in cells], ['6'])

    # -- editable flags -------------------------------------------------------

    def test_control_column_not_editable(self):
        l1 = make_lesson(self.cs_a, D_OCT_A, 1)
        make_control(self.cs_a, date=D_OCT_A, lesson=l1)
        resp = self.client.get(reverse('monthly_journal', kwargs={
            'school_id': self.school_a.id, 'class_name': '7-А', 'subject': MATH,
        }), {'month': OCT})
        idx = col_index(resp.context, D_OCT_A, l1.id)
        self.assertFalse(resp.context['columns'][idx]['editable'])

    def test_locked_quarter_columns_not_editable(self):
        from portal.tests.helpers import lock_quarter
        lock_quarter(self.school_a, '7-А', MATH, quarter=1)
        resp = self.client.get(reverse('monthly_journal', kwargs={
            'school_id': self.school_a.id, 'class_name': '7-А', 'subject': MATH,
        }), {'month': OCT})
        self.assertFalse(resp.context['can_edit'] is False)
        self.assertTrue(all(
            not col['editable'] for col in resp.context['columns']
        ))

    # -- UI contract ----------------------------------------------------------

    def test_quick_pick_markup_present(self):
        resp = self.client.get(reverse('monthly_journal', kwargs={
            'school_id': self.school_a.id, 'class_name': '7-А', 'subject': MATH,
        }), {'month': OCT})
        html = resp.content.decode()
        for n in range(1, 11):
            self.assertIn(f'data-gp="{n}"', html)
        self.assertIn('Қӯшимча', html)
        self.assertIn('leave-modal', html)
        self.assertIn('beforeunload', html)
        self.assertIn('mcell-editable', html)

    def test_daily_quick_pick_markup_present(self):
        resp = self.client.get(reverse('grade_entry', kwargs={
            'school_id': self.school_a.id, 'class_name': '7-А', 'subject': MATH,
        }), {'date': D_OCT_A.isoformat()})
        html = resp.content.decode()
        self.assertIn('quick-pick', html)
        self.assertIn('data-qp', html)
        self.assertIn('showQuickPick', html)
