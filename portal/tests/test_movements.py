"""🔁 Омад-рафт — student mobility journal and А.Ҷ.Ч. export columns.

add_student writes an arrival row, remove_student a departure row that
keeps the name/gender snapshot after the Student record is gone. The
journal page and the zavuch report read the same records.
"""
import io

from django.urls import reverse
from openpyxl import load_workbook

from portal.models import Student, StudentMovement

from .helpers import D_Q1, D_Q1_B, GoldenBase


class MovementRecordingTests(GoldenBase):
    """Automatic rows written by the add/remove views."""

    def test_add_student_records_arrival(self):
        self.client.force_login(self.zavuch_a)
        resp = self.client.post(
            reverse('add_student', args=[self.school_a.id, '7-А']),
            {'full_name': 'Олим Саидов'})
        self.assertEqual(resp.status_code, 302)
        m = StudentMovement.objects.get(school=self.school_a)
        self.assertEqual(m.movement_type, 'arrival')
        self.assertEqual(m.student_name, 'Олим Саидов')
        self.assertEqual(m.class_name, '7-А')
        self.assertEqual(m.gender, 'M')
        self.assertFalse(m.is_manual)
        self.assertEqual(m.created_by, self.zavuch_a)

    def test_remove_student_records_departure_snapshot(self):
        self.assertEqual(self.s2.gender, 'F')   # Мадина — авто-духтар
        self.client.force_login(self.zavuch_a)
        self.client.post(
            reverse('remove_student', args=[self.school_a.id, '7-А']),
            {'student_id': self.s2.id})
        m = StudentMovement.objects.get(school=self.school_a)
        self.assertEqual(m.movement_type, 'departure')
        # Snapshot survives the Student row's deletion.
        self.assertEqual(m.student_name, 'Мадина Каримова')
        self.assertEqual(m.gender, 'F')
        self.assertIsNone(StudentMovement.objects.get().student)

    def test_class_transfer_writes_no_movement(self):
        # Same-school move (7-А -> 7-Б) is not a school-level move.
        self.client.force_login(self.zavuch_a)
        self.client.post(
            reverse('edit_student', args=[self.school_a.id, '7-А']),
            {'student_id': self.s1.id, 'full_name': 'Фирӯз Раҳимов',
             'class_name': '7-Б'})
        self.assertFalse(
            StudentMovement.objects.filter(school=self.school_a).exists())


class MovementJournalTests(GoldenBase):
    """The 🔁 Омад-рафт page: summary grid, record list, manual edits."""

    def _get(self, user, **params):
        self.client.force_login(user)
        return self.client.get(reverse('student_movements'), params)

    def test_zavuch_sees_own_school_summary(self):
        self.client.force_login(self.zavuch_a)
        resp = self._get(self.zavuch_a)
        self.assertEqual(resp.status_code, 200)
        rows = {r['class_name']: r for r in resp.context['summary']['rows']}
        row = rows['7-А']
        # 2 current students, no moves -> аввал=охир=2, 1 духтар.
        self.assertEqual(
            (row['start_total'], row['start_f'],
             row['end_total'], row['end_f']),
            (2, 1, 2, 1))
        # School B's class never leaks in.
        self.assertNotIn('7-Б', rows)
        # Zavuch gets the manual-entry form and delete buttons.
        self.assertContains(resp, 'name="full_name"')

    def test_arrival_and_departure_move_counts(self):
        # s2 leaves (рафт), a new boy arrives (омад) — both inside Q1.
        StudentMovement.objects.create(
            school=self.school_a, student_name='Мадина Каримова',
            class_name='7-А', gender='F', movement_type='departure',
            date=D_Q1, academic_year=D_Q1.year)
        StudentMovement.objects.create(
            school=self.school_a, student_name='Олим Саидов',
            class_name='7-А', gender='M', movement_type='arrival',
            date=D_Q1_B, academic_year=D_Q1.year)
        self.client.force_login(self.zavuch_a)
        resp = self._get(self.zavuch_a, quarter=1)
        row = next(r for r in resp.context['summary']['rows']
                   if r['class_name'] == '7-А')
        # Live roster still 2/1 -> аввал = 2-1+1=2, охир 2; омад 1/0,
        # рафт 1/1.
        self.assertEqual(
            (row['arr_total'], row['arr_f'],
             row['dep_total'], row['dep_f']),
            (1, 0, 1, 1))
        # Other quarters stay empty.
        resp = self._get(self.zavuch_a, quarter=2)
        self.assertEqual(resp.context['summary']['movements'], [])

    def test_teacher_denied(self):
        resp = self._get(self.teacher_a)
        self.assertEqual(resp.status_code, 302)

    def test_principal_reads_journal_but_cannot_edit(self):
        resp = self._get(self.principal_a)
        self.assertEqual(resp.status_code, 200)
        rows = {r['class_name']: r
                for r in resp.context['summary']['rows']}
        self.assertIn('7-А', rows)
        self.assertNotIn('7-Б', rows)
        # Director sees the journal but no manual-entry form.
        self.assertNotContains(resp, 'name="full_name"')
        # Writes stay zavuch-only even inside their own school.
        self.client.post(reverse('movement_add'), {
            'full_name': 'Ҳаким Валиев', 'class_name': '7-А',
            'movement_type': 'arrival', 'date': D_Q1.isoformat()})
        self.assertFalse(StudentMovement.objects.exists())

    def test_staff_sees_picker_but_no_edit(self):
        from .helpers import make_user
        staff = make_user('maorif_staff_mv', staff=True)
        resp = self._get(staff, school_id=self.school_a.id)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'name="school_id"')
        self.assertNotContains(resp, 'name="full_name"')

    def test_manual_add_and_delete(self):
        self.client.force_login(self.zavuch_a)
        resp = self.client.post(reverse('movement_add'), {
            'full_name': 'Ҳаким Валиев',
            'class_name': '7-А',
            'movement_type': 'arrival',
            'date': D_Q1.isoformat(),
            'gender': '',   # авто → писар
            'note': 'Аз мактаби №2',
        })
        self.assertEqual(resp.status_code, 302)
        m = StudentMovement.objects.get(school=self.school_a)
        self.assertTrue(m.is_manual)
        self.assertEqual(m.gender, 'M')
        self.assertEqual(m.note, 'Аз мактаби №2')

        resp = self.client.post(
            reverse('movement_delete'),
            {'movement_id': m.id, 'school_id': self.school_a.id})
        self.assertFalse(StudentMovement.objects.filter(
            school=self.school_a).exists())

    def test_manual_row_edit_prefills_form_and_saves(self):
        m = StudentMovement.objects.create(
            school=self.school_a, student_name='Олим Саидов',
            class_name='7-А', gender='M', movement_type='arrival',
            date=D_Q1, academic_year=D_Q1.year, is_manual=True)
        self.client.force_login(self.zavuch_a)
        # ?edit=N prefills the form and points it at movement_edit.
        resp = self._get(self.zavuch_a, quarter=1, edit=m.id)
        self.assertContains(resp, 'Қайдни таҳрир кардан')
        self.assertContains(resp, 'name="movement_id"')
        self.assertContains(resp, 'value="Олим Саидов"')
        # Saving posts to movement_edit and updates the row in place.
        resp = self.client.post(reverse('movement_edit'), {
            'movement_id': m.id,
            'school_id': self.school_a.id,
            'quarter': 1,
            'full_name': 'Ҳаким Валиев',
            'class_name': '7-А',
            'movement_type': 'departure',
            'date': D_Q1_B.isoformat(),
            'gender': '',
            'note': 'Таҳрир шуд',
        })
        self.assertEqual(resp.status_code, 302)
        m.refresh_from_db()
        self.assertEqual(m.student_name, 'Ҳаким Валиев')
        self.assertEqual(m.movement_type, 'departure')
        self.assertEqual(m.note, 'Таҳрир шуд')
        self.assertEqual(m.gender, 'M')   # авто-аникит → писар
        self.assertTrue(m.is_manual)

    def test_edit_works_on_auto_rows_but_not_other_school(self):
        # Auto rows are editable too — e.g. fixing a departure recorded
        # with the wrong date; the is_manual badge stays as the origin.
        auto = StudentMovement.objects.create(
            school=self.school_a, student_name='Авто Қайд',
            class_name='7-А', gender='M', movement_type='arrival',
            date=D_Q1, academic_year=D_Q1.year, is_manual=False)
        self.client.force_login(self.zavuch_a)
        resp = self._get(self.zavuch_a, quarter=1, edit=auto.id)
        self.assertContains(resp, 'Қайдни таҳрир кардан')
        self.client.post(reverse('movement_edit'), {
            'movement_id': auto.id,
            'school_id': self.school_a.id,
            'quarter': 1,
            'full_name': 'Авто Қайд', 'class_name': '7-А',
            'movement_type': 'departure', 'date': D_Q1_B.isoformat()})
        auto.refresh_from_db()
        self.assertEqual(auto.movement_type, 'departure')
        self.assertEqual(auto.date, D_Q1_B)
        self.assertFalse(auto.is_manual)   # badge keeps the origin

        # A row belonging to another school is never touched.
        other = StudentMovement.objects.create(
            school=self.school_b, student_name='Ҷаҳонгир Назаров',
            class_name='7-Б', gender='M', movement_type='arrival',
            date=D_Q1, academic_year=D_Q1.year, is_manual=True)
        self.client.post(reverse('movement_edit'), {
            'movement_id': other.id,
            'school_id': self.school_a.id,
            'quarter': 1,
            'full_name': 'Тағйир Ёфт', 'class_name': '7-А',
            'movement_type': 'departure', 'date': D_Q1.isoformat()})
        other.refresh_from_db()
        self.assertNotEqual(other.student_name, 'Тағйир Ёфт')

    def test_zavuch_cannot_touch_other_school(self):
        # A manual row for school B exists; school A's zavuch can neither
        # see nor delete it.
        m = StudentMovement.objects.create(
            school=self.school_b, student_name='Ҷаҳонгир Назаров',
            class_name='7-Б', gender='M', movement_type='departure',
            date=D_Q1, academic_year=D_Q1.year)
        self.client.force_login(self.zavuch_a)
        resp = self._get(self.zavuch_a)
        self.assertNotContains(resp, 'Ҷаҳонгир Назаров')
        self.client.post(
            reverse('movement_delete'),
            {'movement_id': m.id, 'school_id': self.school_b.id})
        self.assertTrue(StudentMovement.objects.filter(pk=m.pk).exists())


class MovementExportTests(GoldenBase):
    """А.Ҷ.Ч. Омад/Рафт columns read the journal."""

    def _ws(self, quarter=1):
        self.client.force_login(self.zavuch_a)
        resp = self.client.get(
            reverse('zavuch_report_export'), {'quarter': quarter})
        return load_workbook(io.BytesIO(resp.content)).active

    def test_movements_fill_omad_raft_columns(self):
        # s2 (духтар) leaves 7-А; Олим (писар) joins 7-А — Q1 dates.
        StudentMovement.objects.create(
            school=self.school_a, student_name='Мадина Каримова',
            class_name='7-А', gender='F', movement_type='departure',
            date=D_Q1, academic_year=D_Q1.year)
        StudentMovement.objects.create(
            school=self.school_a, student_name='Олим Саидов',
            class_name='7-А', gender='M', movement_type='arrival',
            date=D_Q1_B, academic_year=D_Q1.year)
        r = 10   # СИНФИ 7 row (FIRST_DATA_ROW + 7 - 2)
        ws = self._ws(quarter=1)
        # Live roster 7-А: 2 total / 1 girl. Аввал ҳамагӣ = 2-1+1 = 2;
        # аввал духтарон = 1-0+1 = 2 (the departed girl was enrolled at
        # quarter start). Омад 1/0, рафт 1/1.
        self.assertEqual(
            [ws[f'{c}{r}'].value for c in 'CDEFGH'],
            [2, 2, 1, 0, 1, 1])
        # Q2 has no movements -> blank counts and C = live roster.
        ws = self._ws(quarter=2)
        self.assertEqual(
            [ws[f'{c}{r}'].value for c in 'CDEFGH'],
            [2, 1, 0, 0, 0, 0])


class MovementListExportTests(GoldenBase):
    """The official «Рӯйхати хонандагони…» header and the Excel export."""

    def test_page_shows_official_list_title(self):
        from portal.views import academic_year_start
        self.client.force_login(self.zavuch_a)
        resp = self.client.get(reverse('student_movements'), {'quarter': 1})
        ay = academic_year_start()
        self.assertContains(
            resp,
            'Рӯйхати хонандагони ба муассиса қабулшуда ва аз муассиса '
            f'хориҷшуда дар соли таҳсили {ay}–{ay + 1} '
            f'дар {self.school_a.name}')

    def test_excel_export_carries_header_and_rows(self):
        StudentMovement.objects.create(
            school=self.school_a, student_name='Мадина Каримова',
            class_name='7-А', gender='F', movement_type='departure',
            date=D_Q1, academic_year=D_Q1.year, note='Ҳуҷҷат №5')
        self.client.force_login(self.zavuch_a)
        resp = self.client.get(reverse('movements_export'), {'quarter': 1})
        self.assertEqual(resp.status_code, 200)
        self.assertIn('attachment', resp['Content-Disposition'])
        ws = load_workbook(io.BytesIO(resp.content)).active
        self.assertIn('Рӯйхати хонандагони ба муассиса қабулшуда',
                      ws['A1'].value)
        self.assertIn(self.school_a.name, ws['A1'].value)
        self.assertEqual(
            [ws.cell(row=3, column=c).value for c in range(1, 8)],
            ['№', 'Сана', 'Ному насаб', 'Синф', 'Навъ', 'Ҷинс', 'Эзоҳ'])
        self.assertEqual(
            [ws.cell(row=4, column=c).value for c in range(1, 8)],
            [1, D_Q1.strftime('%d.%m.%Y'), 'Мадина Каримова',
             '7-А', 'Рафт', 'Духтар', 'Ҳуҷҷат №5'])

    def test_excel_export_scoped_to_school_and_quarter(self):
        # Other school's row and another quarter's row stay out.
        StudentMovement.objects.create(
            school=self.school_b, student_name='Ҷаҳонгир Назаров',
            class_name='7-Б', gender='M', movement_type='arrival',
            date=D_Q1, academic_year=D_Q1.year)
        self.client.force_login(self.zavuch_a)
        resp = self.client.get(reverse('movements_export'), {'quarter': 2})
        self.assertEqual(resp.status_code, 200)
        ws = load_workbook(io.BytesIO(resp.content)).active
        self.assertIn('сабт нашудааст', ws['A4'].value)

    def test_teacher_denied_export(self):
        self.client.force_login(self.teacher_a)
        resp = self.client.get(reverse('movements_export'), {'quarter': 1})
        self.assertEqual(resp.status_code, 302)
