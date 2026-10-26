"""📁 Солҳои таҳсил — quarter closeout, archive, and reopen access.

GoldenBase gives School A (7-А graded, 1-А sticker class territory via
students only) and School B (7-Б). The zavuch closes quarters for their
own school; маориф/админ reopen.
"""
import datetime
import io
import tempfile

from django.test import override_settings
from django.urls import reverse
from openpyxl import load_workbook

from portal.models import QuarterLock
from portal.quarter_close import (archive_path, collect_quarter_close,
                                  save_quarter_archive)

from .helpers import (GoldenBase, MATH, TAJIK, make_class_subject,
                      make_grade, make_quarter_grade, make_user)


_TMP_MEDIA = tempfile.mkdtemp(prefix='qc_archive_')


@override_settings(MEDIA_ROOT=_TMP_MEDIA)
class QuarterCloseTests(GoldenBase):

    def setUp(self):
        super().setUp()
        # School A gets a second graded subject; the auto-seeded
        # curriculum subjects stay active too.
        make_class_subject(self.school_a, '7-А', TAJIK,
                           teacher=self.teacher_a)
        from portal.views import academic_year_start
        self.ay = academic_year_start()

    # ---------------------------------------------------------- status

    def test_status_counts_locks_and_missing(self):
        qs = collect_quarter_close(self.school_a, 1, self.ay)
        self.assertFalse(qs['closed'])
        self.assertGreater(qs['total'], 0)
        self.assertEqual(qs['locked_count'], 0)
        # No QuarterGrade rows yet -> every student counts as missing.
        self.assertGreater(qs['missing_total'], 0)

        make_quarter_grade(self.s1, quarter=1, grade=9)
        make_quarter_grade(self.s2, quarter=1, grade=8)
        qs = collect_quarter_close(self.school_a, 1, self.ay)
        # МАТЕМАТИКА now fully graded — its missing drops to 0, other
        # subjects keep their gaps.
        math_row = next(r for r in qs['rows']
                        if r['subject'] == MATH)
        self.assertEqual(math_row['missing'], 0)

    # ---------------------------------------------------------- close

    def test_close_requires_tasdiq_word(self):
        self.client.force_login(self.zavuch_a)
        resp = self.client.post(reverse('quarter_close'), {
            'school_id': self.school_a.id, 'quarter': 1, 'year': self.ay,
            'confirm': 'ҳа'})
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(QuarterLock.objects.filter(
            school=self.school_a).exists())

        resp = self.client.post(reverse('quarter_close'), {
            'school_id': self.school_a.id, 'quarter': 1, 'year': self.ay,
            'confirm': 'ТАСДИҚ'})
        self.assertEqual(resp.status_code, 302)
        qs = collect_quarter_close(self.school_a, 1, self.ay)
        self.assertTrue(qs['closed'])
        self.assertEqual(qs['locked_count'], qs['total'])
        # Other quarters and school B stay untouched.
        self.assertFalse(
            collect_quarter_close(self.school_a, 2, self.ay)['closed'])
        self.assertFalse(
            collect_quarter_close(self.school_b, 1, self.ay)['closed'])

    def test_close_locks_grades(self):
        self.client.force_login(self.zavuch_a)
        self.client.post(reverse('quarter_close'), {
            'school_id': self.school_a.id, 'quarter': 1,
            'year': self.ay, 'confirm': 'ТАСДИҚ'})
        self.client.force_login(self.teacher_a)
        resp = self.client.post(reverse('save_grade_ajax'), {
            'student_id': self.s1.id, 'subject': MATH,
            'grade_type': 'daily', 'score': '8',
            'date': f'{self.ay}-10-15',
        })
        self.assertIn('баста', resp.json().get('message', '').lower())

    def test_zavuch_cannot_close_other_school(self):
        self.client.force_login(self.zavuch_a)
        self.client.post(reverse('quarter_close'), {
            'school_id': self.school_b.id, 'quarter': 1,
            'year': self.ay, 'confirm': 'ТАСДИҚ'})
        self.assertFalse(QuarterLock.objects.filter(
            school=self.school_b).exists())

    def test_teacher_cannot_close(self):
        self.client.force_login(self.teacher_a)
        self.client.post(reverse('quarter_close'), {
            'school_id': self.school_a.id, 'quarter': 1,
            'year': self.ay, 'confirm': 'ТАСДИҚ'})
        self.assertFalse(QuarterLock.objects.filter(
            school=self.school_a).exists())

    # ---------------------------------------------------------- reopen

    def _close_q1(self):
        self.client.force_login(self.zavuch_a)
        self.client.post(reverse('quarter_close'), {
            'school_id': self.school_a.id, 'quarter': 1,
            'year': self.ay, 'confirm': 'ТАСДИҚ'})

    def test_reopen_only_maorif(self):
        self._close_q1()
        # Zavuch cannot reopen their own closed quarter.
        resp = self.client.post(reverse('quarter_reopen'), {
            'school_id': self.school_a.id, 'quarter': 1,
            'year': self.ay})
        self.assertTrue(collect_quarter_close(
            self.school_a, 1, self.ay)['closed'])

        self.client.force_login(self.director)
        self.client.post(reverse('quarter_reopen'), {
            'school_id': self.school_a.id, 'quarter': 1,
            'year': self.ay})
        qs = collect_quarter_close(self.school_a, 1, self.ay)
        self.assertFalse(qs['closed'])
        self.assertEqual(qs['locked_count'], 0)

    # ---------------------------------------------------------- page

    def test_years_page_shows_quarters(self):
        self.client.force_login(self.zavuch_a)
        resp = self.client.get(reverse('academic_years'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Соли таҳсили')
        self.assertContains(resp, 'Чоряки')

    # ---------------------------------------------------------- archive

    def test_archive_workbook_contents(self):
        make_quarter_grade(self.s1, subject=MATH, quarter=1, grade=9)
        make_quarter_grade(self.s2, subject=TAJIK, quarter=1, grade=7)
        make_grade(self.s1, attendance='-', date=datetime.date(
            self.ay, 10, 10))
        buf = save_quarter_archive(self.school_a, 1, self.ay)
        self.assertTrue(buf.exists())
        wb = load_workbook(io.BytesIO(buf.read_bytes()))
        self.assertIn('7-А', wb.sheetnames)
        ws = wb['7-А']
        self.assertEqual(ws.cell(2, 1).value, '№')
        header = [c.value for c in ws[2]]
        self.assertIn('Хонанда', header)
        self.assertIn(MATH, header)
        # Grade written to the right subject cell.
        name_col = header.index('Хонанда') + 1
        math_col = header.index(MATH) + 1
        row = None
        for r in range(3, ws.max_row + 1):
            if ws.cell(r, name_col).value == 'Фирӯз Раҳимов':
                row = r
        self.assertIsNotNone(row)
        self.assertEqual(ws.cell(row, math_col).value, 9)
        # Attendance sheet has the unexcused mark.
        self.assertIn('Давомат', wb.sheetnames)
        att = wb['Давомат']
        found = any(att.cell(r, 2).value == 'Фирӯз Раҳимов'
                    and att.cell(r, 4).value == 1
                    for r in range(3, att.max_row + 1))
        self.assertTrue(found)

    def test_archive_download(self):
        self._close_q1()
        resp = self.client.get(reverse('quarter_archive_download'), {
            'quarter': 1, 'year': self.ay})
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheetml', resp['Content-Type'])
        self.assertTrue(archive_path(
            self.school_a, 1, self.ay).exists())
