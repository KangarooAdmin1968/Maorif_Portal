"""Ҳуҷҷатнигорӣ → Ҳисоботи завуч — А.Ҷ.Ч. quarterly summary Excel export.

Reproduces the owner's workbook (portal/excel_templates/zavuch_report.xlsx):
one row per grade level 2-11 with parallel classes aggregated, the
header block, band formulas and signature lines preserved.

Data sources (existing portal data only):
  - Student (class level, gender)               -> C, D (+Эзоҳ row for level 1)
  - QuarterGrade.grade for the chosen quarter    -> K, L and band cells N/P/T
  - StudentMovement (Омад/Рафт journal)          -> E..H, and the
    quarter-start baseline inside C/D
  - School.name / School.director                -> title + signature
"""
import io
from collections import defaultdict
from pathlib import Path

from django.conf import settings
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from .class_report import QUARTER_ORDINALS, _director_name
from .models import QuarterGrade, Student, StudentMovement
from .utils import class_numeric_part

TEMPLATE_PATH = (
    Path(settings.BASE_DIR) / 'portal' / 'excel_templates' / 'zavuch_report.xlsx'
)

FIRST_DATA_ROW = 5
LEVELS = range(2, 12)   # report rows map to grade levels 2..11
NOTE_ROW = 17
DIRECTOR_ROW = 20
DEPUTY_ROW = 21

# Band lower bounds for a student's quarter average: >=9 -> 9-10 (N),
# >=7 -> 7-8 (P), 4-6 is the remainder formula (R), <4 -> 1-3 (T).
BAND_9_10, BAND_7_8, BAND_4_6, BAND_1_3 = 'hi', 'md', 'rm', 'lo'


class ZavuchReportError(Exception):
    """Raised when the school summary cannot be rendered."""


def _band(avg):
    if avg >= 9:
        return BAND_9_10
    if avg >= 7:
        return BAND_7_8
    if avg >= 4:
        return BAND_4_6
    return BAND_1_3


def _whole_grade(avg):
    """Round-half-up the quarter average to the nearest mark 1..10 —
    feeds the per-grade Хол columns of the АЧЧ school report and the
    oblast maorif variant."""
    return max(1, min(10, int(avg + 0.5)))


def _pct(part, whole):
    """Rounded part/whole*100 as a plain value — displays even in
    Excel's Protected View where formulas are never recalculated."""
    return round(part / whole * 100, 2) if whole else 0


def _collect_level_stats(school, quarter, ay, windows):
    """Per-level roster / movement / quarter-grade aggregates.

    Returns (levels, class_counts, moved, graded) keyed by grade level
    1..11. ``levels`` = live roster [total, girls]; ``class_counts`` =
    parallel-class count; ``moved`` = quarter-window mobility
    [arr, arr_f, dep, dep_f]; ``graded`` = {'total', 'f', 'bands', 'dist'}
    where dist maps the rounded per-student average 1..10 to a count.
    """
    levels = {lvl: [0, 0] for lvl in range(1, 12)}   # [total, girls]
    class_counts = defaultdict(set)
    for s in Student.objects.filter(school=school).values(
            'class_name', 'gender'):
        lvl = class_numeric_part(s['class_name'])
        if lvl in levels:
            levels[lvl][0] += 1
            class_counts[lvl].add(s['class_name'])
            if s['gender'] == 'F':
                levels[lvl][1] += 1

    qstart, qend = windows[quarter]
    moved = {lvl: [0, 0, 0, 0] for lvl in range(1, 12)}
    for m in StudentMovement.objects.filter(
            school=school, date__gte=qstart, date__lte=qend).values(
            'class_name', 'gender', 'movement_type'):
        lvl = class_numeric_part(m['class_name'])
        if lvl not in moved:
            continue
        i = 0 if m['movement_type'] == 'arrival' else 2
        moved[lvl][i] += 1
        if m['gender'] == 'F':
            moved[lvl][i + 1] += 1

    # Quarter grades -> per-student average -> band + whole mark,
    # grouped by level.
    sums = {}
    for row in QuarterGrade.objects.filter(
            student__school=school, quarter=quarter,
            academic_year=ay, grade__isnull=False).values(
            'student_id', 'student__class_name', 'student__gender',
            'grade'):
        lvl = class_numeric_part(row['student__class_name'])
        if lvl not in levels:
            continue
        entry = sums.setdefault(
            row['student_id'],
            [0.0, 0, lvl, row['student__gender']])
        entry[0] += float(row['grade'])
        entry[1] += 1

    graded = {lvl: {'total': 0, 'f': 0,
                    'bands': defaultdict(int),
                    'dist': defaultdict(int)}
              for lvl in range(1, 12)}
    for total, count, lvl, gender in sums.values():
        g = graded[lvl]
        g['total'] += 1
        if gender == 'F':
            g['f'] += 1
        avg = total / count
        g['bands'][_band(avg)] += 1
        g['dist'][_whole_grade(avg)] += 1

    return levels, class_counts, moved, graded


def build_zavuch_report_workbook(school, quarter):
    """Fill the owner's А.Ҷ.Ч. template for one school + quarter.

    Returns a BytesIO xlsx. Per-student band = average of the student's
    QuarterGrade.grade rows in the chosen quarter/current academic year;
    students without any quarter grade land in Бе хол (V = I-K).
    """
    from .views import (  # local imports avoid cycles
        academic_year_start, academic_year_quarter_windows)

    if quarter not in (1, 2, 3, 4):
        raise ZavuchReportError('Чораки нодуруст.')

    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active
    ay = academic_year_start()
    levels, _cc, moved, graded = _collect_level_stats(
        school, quarter, ay, academic_year_quarter_windows(ay))

    # Header wording: the template's «Шумораи» clips to «мор» in narrow
    # columns — write the official «Шумора» on every count column.
    for col in 'NPRTV':
        ws[f'{col}4'] = 'Шумора'

    for i, lvl in enumerate(LEVELS):
        r = FIRST_DATA_ROW + i
        ws[f'A{r}'] = i + 1
        ws[f'B{r}'] = f'СИНФИ {lvl}'
        arr_t, arr_f, dep_t, dep_f = moved[lvl]
        ws[f'C{r}'] = levels[lvl][0] - arr_t + dep_t
        ws[f'D{r}'] = levels[lvl][1] - arr_f + dep_f
        ws[f'E{r}'] = arr_t
        ws[f'F{r}'] = arr_f
        ws[f'G{r}'] = dep_t
        ws[f'H{r}'] = dep_f
        ws[f'I{r}'] = f'=C{r}+E{r}-G{r}'
        ws[f'J{r}'] = f'=D{r}+F{r}-H{r}'
        g = graded[lvl]
        k = g['total']
        n, p, t = (g['bands'][BAND_9_10], g['bands'][BAND_7_8],
                   g['bands'][BAND_1_3])
        r46 = k - n - p - t   # 4-6 remainder — owner's K-N-P-T logic
        ws[f'K{r}'] = k
        ws[f'L{r}'] = g['f']
        ws[f'M{r}'] = 100
        # Band cells are computed values (nothing manual affects them);
        # only the Омад/Рафт-dependent chain stays formula-driven so
        # hand-entered E..H numbers fold into I/J/V/W on recalc.
        ws[f'N{r}'] = n
        ws[f'O{r}'] = _pct(n, k)
        ws[f'P{r}'] = p
        ws[f'Q{r}'] = _pct(p, k)
        ws[f'R{r}'] = r46
        ws[f'S{r}'] = _pct(r46, k)
        ws[f'T{r}'] = t
        ws[f'U{r}'] = _pct(t, k)
        ws[f'V{r}'] = f'=I{r}-K{r}'
        ws[f'W{r}'] = f'=IF($K{r}=0,0,V{r}*100/$K{r})'
        ws[f'X{r}'] = _pct(n + p, k)
        ws[f'Y{r}'] = 100

    ordinal = QUARTER_ORDINALS.get(quarter, f'{quarter}-ум')
    ws['A1'] = (f'Натиҷаи ҷамъбастии чоряк (А.Ҷ.Ч.) {ordinal} '
                f'дар {school.name}\nдар соли таҳсили {ay}-{ay + 1}')
    n1, f1 = levels[1]
    ws[f'A{NOTE_ROW}'] = (
        'Эзоҳ: Шумораи хонандагони синфҳои 1-ум '
        f'{n1} нафар аз ҷумла духтарон {f1} нафарро ташкил медиҳад.')
    director = _director_name(school)
    ws[f'A{DIRECTOR_ROW}'] = (
        f'Директори муассиса:    {director}'.rstrip())
    # Deputy name has no reliable portal source — left blank by design.
    ws[f'A{DEPUTY_ROW}'] = 'Муовини директор оид ба таълим:'

    wb.calculation.fullCalcOnLoad = True
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# ----------------------------------------------------------------------
# АЧЧ variant: every mark 10..1 gets its own count + % column, rows are
# grade levels (only those present) plus group subtotals (1-4 / 5-9 /
# 10-11) and a grand total — the official school quarterly form.
# ----------------------------------------------------------------------

GRADES_LAST_COL = 36       # A..AJ
GRADES_FIRST_DATA_ROW = 5
GRADES_GROUPS = ((1, 4), (5, 9), (10, 11))   # paper groupings

_thin = Side(style='thin')
_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_CENTER = Alignment(horizontal='center', vertical='center',
                    wrap_text=True)
_LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)


def _grades_header(ws, quarter):
    """Rows 2-4 of the АЧЧ form: group captions + sub-headers."""

    def cell(r, c, text, size=9, bold=True):
        x = ws.cell(row=r, column=c, value=text)
        x.font = Font(bold=bold, size=size)
        x.alignment = _CENTER

    def merge(text, col, r1=2, r2=4, c2=None):
        ws.merge_cells(start_row=r1, start_column=col,
                       end_row=r2, end_column=c2 or col)
        cell(r1, col, text)

    merge('№', 1)
    merge('Синфҳо', 2)
    merge('Миқдори синфҳо', 3)
    # Paired roster columns: group caption on row 2, Ҳамагӣ / аз ҷумла,
    # духтарон merged down rows 3-4.
    pairs = (
        (4, f'Шумораи хонандагон дар аввали чоряки {quarter}'),
        (6, 'Омад'),
        (8, 'Рафт'),
        (10, f'Шумораи хонандагон дар охири чоряки {quarter}'),
        (12, 'Баҳо гирифтан'),
    )
    for col, text in pairs:
        ws.merge_cells(start_row=2, start_column=col,
                       end_row=2, end_column=col + 1)
        cell(2, col, text)
        for sub, label in ((col, 'Ҳамагӣ'),
                           (col + 1, 'аз ҷумла, духтарон')):
            ws.merge_cells(start_row=3, start_column=sub,
                           end_row=4, end_column=sub)
            cell(3, sub, label)
    merge('%-и баҳо гирифтани', 14)
    ws.merge_cells(start_row=2, start_column=15,
                   end_row=2, end_column=34)
    cell(2, 15, 'Баҳо')
    col = 15
    for g in range(10, 0, -1):
        ws.merge_cells(start_row=3, start_column=col,
                       end_row=3, end_column=col + 1)
        cell(3, col, str(g))
        cell(4, col, 'Шумора')
        cell(4, col + 1, '%')
        col += 2
    merge('Сифати таҳсилот % (холҳои 7-10)', 35)
    merge('Азхудкунии барномаҳои таълимӣ % (холҳои 4-10)', 36)
    ws.row_dimensions[2].height = 34
    ws.row_dimensions[3].height = 46
    ws.row_dimensions[4].height = 26


def _grades_row(ws, r, index, name, data, bold=False):
    """One АЧЧ table row. `data` carries roster/movement/graded fields."""
    if index is not None:
        ws.cell(row=r, column=1, value=index)
    name_cell = ws.cell(row=r, column=2, value=name)
    name_cell.alignment = _LEFT
    values = (
        data['classes'],
        data['start'], data['start_f'],
        data['arr'], data['arr_f'],
        data['dep'], data['dep_f'],
        data['end'], data['end_f'],
        data['graded'], data['graded_f'],
        round(_pct(data['graded'], data['end'])),
    )
    for col, value in zip(range(3, 15), values):
        c = ws.cell(row=r, column=col, value=value)
        if bold:
            c.font = Font(bold=True)
    col = 15
    g = data['graded']
    for mark in range(10, 0, -1):
        n = data['dist'][mark]
        c = ws.cell(row=r, column=col, value=n)
        if bold:
            c.font = Font(bold=True)
        c = ws.cell(row=r, column=col + 1,
                    value=round(_pct(n, g)))
        if bold:
            c.font = Font(bold=True)
        col += 2
    dist = data['dist']
    quality = sum(dist[m] for m in range(7, 11))
    mastery = g - sum(dist[m] for m in range(1, 4))
    for col, value in ((35, round(_pct(quality, g))),
                       (36, round(_pct(mastery, g)))):
        c = ws.cell(row=r, column=col, value=value)
        if bold:
            c.font = Font(bold=True)
    if bold:
        name_cell.font = Font(bold=True)


def _level_data(lvl, levels, class_counts, moved, graded):
    arr_t, arr_f, dep_t, dep_f = moved[lvl]
    end, end_f = levels[lvl]
    g = graded[lvl]
    return {
        'classes': len(class_counts[lvl]),
        'start': end - arr_t + dep_t, 'start_f': end_f - arr_f + dep_f,
        'arr': arr_t, 'arr_f': arr_f,
        'dep': dep_t, 'dep_f': dep_f,
        'end': end, 'end_f': end_f,
        'graded': g['total'], 'graded_f': g['f'],
        'dist': g['dist'],
    }


def _sum_data(items):
    """Group subtotal / grand total: counts add, dist merges."""
    out = {k: sum(i[k] for i in items)
           for k in ('classes', 'start', 'start_f', 'arr', 'arr_f',
                     'dep', 'dep_f', 'end', 'end_f',
                     'graded', 'graded_f')}
    out['dist'] = defaultdict(int)
    for i in items:
        for m, n in i['dist'].items():
            out['dist'][m] += n
    return out


def build_zavuch_grades_workbook(school, quarter):
    """АЧЧ school form — one row per grade level present, group
    subtotals (1-4 / 5-9 / 10-11) and a grand total -> BytesIO xlsx."""
    from .views import (  # local imports avoid cycles
        academic_year_start, academic_year_quarter_windows)

    if quarter not in (1, 2, 3, 4):
        raise ZavuchReportError('Чораки нодуруст.')

    ay = academic_year_start()
    levels, class_counts, moved, graded = _collect_level_stats(
        school, quarter, ay, academic_year_quarter_windows(ay))

    wb = Workbook()
    ws = wb.active
    ws.title = f'АЧЧ-{quarter}'

    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1, end_column=GRADES_LAST_COL)
    ws['A1'] = (f'Натиҷаи ҷамъбастии чоряки {quarter} (АЧЧ{quarter}) '
                f'дар {school.name} дар соли таҳсили {ay}-{ay + 1}')
    ws['A1'].font = Font(bold=True, size=12)
    ws['A1'].alignment = _CENTER

    _grades_header(ws, quarter)

    r = GRADES_FIRST_DATA_ROW
    index = 0
    present = [lvl for lvl in range(1, 12) if class_counts[lvl]]
    all_data = []
    for lo, hi in GRADES_GROUPS:
        band = [lvl for lvl in present if lo <= lvl <= hi]
        if not band:
            continue    # e.g. лицей has no 1-4 — skip the whole group
        band_data = []
        for lvl in band:
            index += 1
            data = _level_data(lvl, levels, class_counts, moved, graded)
            _grades_row(ws, r, index, f'Синфи {lvl}', data)
            band_data.append(data)
            r += 1
        sub = _sum_data(band_data)
        _grades_row(ws, r, None, f'Ҳамагӣ (синфҳои {lo}-{hi})', sub,
                    bold=True)
        all_data.extend(band_data)
        r += 1
    if all_data:
        lo, hi = present[0], present[-1]
        label = (f'Ҳамагӣ (синфи {lo})' if lo == hi
                 else f'Ҳамагӣ (синфҳои {lo}-{hi})')
        _grades_row(ws, r, None, label, _sum_data(all_data), bold=True)
        r += 1

    # Borders over the whole table incl. the three header rows.
    for row in range(2, r):
        for col in range(1, GRADES_LAST_COL + 1):
            c = ws.cell(row=row, column=col)
            c.border = _BORDER
            if c.alignment.horizontal is None:
                c.alignment = _CENTER

    sig = r + 1
    ws.cell(row=sig, column=2, value='Директор').font = Font(bold=True)
    ws.cell(row=sig, column=20, value=_director_name(school)).font = (
        Font(bold=True, underline='single'))

    widths = {1: 4, 2: 16, 14: 9, 35: 10, 36: 12}
    for col in range(1, GRADES_LAST_COL + 1):
        ws.column_dimensions[get_column_letter(col)].width = (
            widths.get(col, 6))
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4

    wb.calculation.fullCalcOnLoad = True
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
